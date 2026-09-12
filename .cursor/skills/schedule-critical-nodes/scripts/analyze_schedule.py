#!/usr/bin/env python3
"""OPTIONAL draft analyzer: keyword/template → critical_nodes.

Not the primary path. Agent should reason over the schedule and write via
`write_critical_nodes.py`. Use this only for a rough draft to review.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

# ---------------------------------------------------------------------------
# Headers / keywords (aligned with matching-rules.md + Skill 1 schedule_csv)
# ---------------------------------------------------------------------------

_NAME_HEADERS = {
    "里程碑",
    "节点",
    "节点名称",
    "任务名称",
    "任务",
    "活动",
    "activity",
    "task",
    "task name",
    "task_name",
    "name",
    "milestone",
    "milestone name",
    "description",
    "wbs",
    "wbs name",
}
_START_HEADERS = {
    "开始",
    "开始日期",
    "start",
    "start date",
    "start_date",
    "计划开始",
    "基准开始",
    "日期",
    "date",
    "基准日期",
}
_END_HEADERS = {
    "结束",
    "结束日期",
    "finish",
    "finish date",
    "finish_date",
    "计划完成",
    "完成日期",
    "end",
    "end date",
    "end_date",
    "计划结束",
}
_PACKAGE_HEADERS = {
    "采购包",
    "包名",
    "package",
    "package name",
    "package_name",
    "专业",
    "discipline",
    "trade",
}

_DATE_PATTERNS = (
    ("%Y-%m-%d", r"^\d{4}-\d{2}-\d{2}$"),
    ("%Y/%m/%d", r"^\d{4}/\d{1,2}/\d{1,2}$"),
    ("%Y.%m.%d", r"^\d{4}\.\d{1,2}\.\d{1,2}$"),
    ("%Y%m%d", r"^\d{8}$"),
)

NODE_LABELS: dict[str, tuple[str, str]] = {
    "project_start": ("开工", "Commencement / Start"),
    "project_interface": ("关键界面", "Key Interface"),
    "project_completion": ("竣工", "Practical Completion"),
    "project_handover": ("移交 / 开业", "Handover / Opening"),
    "po_confirm": ("下单确认", "PO Confirmed"),
    "fat": ("出厂 / 厂验", "FAT / Ex-works"),
    "shipment": ("发运", "Shipment"),
    "delivery": ("到货验收", "Delivery / Arrival Acceptance"),
    "site_access": ("进场", "Site Access / Mobilization"),
    "install_complete": ("安装完成", "Installation Complete"),
    "commissioning": ("调试", "Commissioning"),
    "prelim_accept": ("初步验收", "Preliminary Acceptance"),
    "final_accept": ("最终验收", "Final Acceptance"),
    "warranty_start": ("质保起算", "Warranty Start"),
    "contract_sign": ("签约", "Contract Signed"),
    "concept": ("方案", "Concept / Scheme"),
    "detail_deliverable": ("施工图 / 成果交付", "Detail Deliverable"),
    "site_support": ("现场配合", "Site Support"),
    "deliverable_accept": ("成果验收", "Deliverable Acceptance"),
    "gc_milestone": ("总包里程碑", "GC Milestone"),
}

NODE_KEYWORDS: dict[str, list[str]] = {
    "project_start": ["开工", "commencement", "project start"],
    "project_interface": ["界面", "interface", "结构移交", "精装进场"],
    "project_completion": ["竣工", "practical completion", "pc date", "完工"],
    "project_handover": ["移交", "handover", "开业", "opening", "soft opening"],
    "po_confirm": ["下单", "po confirm", "order confirm", "签约下单"],
    "fat": ["厂验", "fat", "factory acceptance", "出厂检验", "ex-works"],
    "shipment": ["发运", "shipment", "shipping", "出运", "etd", "eta"],
    "delivery": ["到货", "delivery", "arrival", "材料到场", "material delivery", "送货"],
    "site_access": ["进场", "site access", "mobilization", "进场安装"],
    "install_complete": [
        "安装完成",
        "installation complete",
        "install finish",
        "安装结束",
    ],
    "commissioning": ["调试", "commissioning", "联调", "testing & commissioning"],
    "prelim_accept": ["初步验收", "preliminary acceptance", "初验", "snagging"],
    "final_accept": ["最终验收", "final acceptance", "终验", "taking over"],
    "warranty_start": ["质保起算", "warranty start", "dlp start"],
    "concept": ["方案", "concept", "schematic", "概念设计"],
    "detail_deliverable": [
        "施工图",
        "detail design",
        "成果交付",
        "deliverable",
        "ifc",
    ],
    "site_support": ["现场配合", "site support", "驻场"],
    "deliverable_accept": ["成果验收", "deliverable accept", "终稿验收"],
}

# Default nodes to expand per package type (MVP)
DEFAULT_NODES: dict[str, list[str]] = {
    "供应": ["delivery"],
    "供装": [
        "delivery",
        "site_access",
        "install_complete",
        "prelim_accept",
        "final_accept",
    ],
    "顾问": ["concept", "detail_deliverable", "deliverable_accept"],
    "总包": [],  # only project cross-link + explicit gc rows
    "_project": ["project_start", "project_completion", "project_handover"],
}

# Fraction windows within coarse bar [0,1] → (start_frac, end_frac)
INFER_FRACTIONS: dict[str, dict[str, tuple[float, float]]] = {
    "供应": {"delivery": (0.0, 1.0)},
    "供装": {
        "delivery": (0.0, 0.25),
        "site_access": (0.20, 0.35),
        "install_complete": (0.35, 0.75),
        "prelim_accept": (0.75, 0.90),
        "final_accept": (0.90, 1.0),
    },
    "顾问": {
        "concept": (0.0, 0.35),
        "detail_deliverable": (0.35, 0.80),
        "deliverable_accept": (0.80, 1.0),
    },
}

NODE_ORDER: dict[str, int] = {
    code: i
    for i, code in enumerate(
        [
            "project_start",
            "project_interface",
            "po_confirm",
            "contract_sign",
            "fat",
            "shipment",
            "delivery",
            "site_access",
            "concept",
            "detail_deliverable",
            "site_support",
            "install_complete",
            "commissioning",
            "prelim_accept",
            "deliverable_accept",
            "final_accept",
            "project_completion",
            "project_handover",
            "warranty_start",
            "gc_milestone",
        ]
    )
}


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def _normalize_header(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip().lower())


def _header_key(header: str) -> str | None:
    norm = _normalize_header(header)
    if not norm:
        return None
    if norm in _NAME_HEADERS:
        return "name"
    if norm in _START_HEADERS:
        return "start"
    if norm in _END_HEADERS:
        return "end"
    if norm in _PACKAGE_HEADERS:
        return "package"
    return None


def parse_date_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        try:
            return f"{int(value.year):04d}-{int(value.month):02d}-{int(value.day):02d}"
        except (TypeError, ValueError):
            pass
    text = str(value).strip()
    if not text:
        return ""
    for fmt, pattern in _DATE_PATTERNS:
        if re.match(pattern, text):
            try:
                return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
    match = re.search(r"\d{4}-\d{2}-\d{2}", text)
    if match:
        return match.group(0)
    match = re.search(r"\d{4}/\d{1,2}/\d{1,2}", text)
    if match:
        try:
            parts = match.group(0).split("/")
            return datetime(int(parts[0]), int(parts[1]), int(parts[2])).strftime(
                "%Y-%m-%d"
            )
        except ValueError:
            pass
    return ""


def _compose_name(base_name: str, package_name: str) -> str:
    base = (base_name or "").strip()
    package = (package_name or "").strip()
    if not base:
        return ""
    if not package:
        return base
    if package.lower() in base.lower():
        return base
    return f"{package} {base}"


def _parse_rows(
    fieldnames: list[str], rows: list[dict[str, Any]]
) -> list[dict[str, str]]:
    mapping: dict[str, str] = {}
    for header in fieldnames:
        key = _header_key(str(header))
        if key and key not in mapping:
            mapping[key] = str(header)

    if "name" not in mapping:
        raise ValueError(
            f"Schedule missing name column. Found headers: {fieldnames}"
        )
    if "start" not in mapping and "end" not in mapping:
        raise ValueError(
            f"Schedule missing date column. Found headers: {fieldnames}"
        )

    milestones: list[dict[str, str]] = []
    for row in rows:
        if not row:
            continue
        name = str(row.get(mapping["name"], "") or "").strip()
        if not name or name.startswith("#"):
            continue
        package = ""
        if mapping.get("package"):
            package = str(row.get(mapping["package"], "") or "").strip()
        display = _compose_name(name, package)
        start = (
            parse_date_cell(row.get(mapping["start"], ""))
            if mapping.get("start")
            else ""
        )
        end = (
            parse_date_cell(row.get(mapping["end"], ""))
            if mapping.get("end")
            else ""
        )
        if not start and end:
            start = end
        if start and not end:
            end = start
        if not start:
            continue
        item: dict[str, str] = {
            "name": display,
            "date": start,
            "raw_name": name,
            "package_hint": package,
        }
        if end and end != start:
            item["end_date"] = end
        else:
            item["end_date"] = start
        milestones.append(item)
    return milestones


def parse_schedule_csv(path: Path) -> list[dict[str, str]]:
    raw = path.read_text(encoding="utf-8-sig")
    reader = csv.DictReader(io.StringIO(raw))
    if not reader.fieldnames:
        return []
    rows = [dict(r) for r in reader]
    return _parse_rows(list(reader.fieldnames), rows)


def parse_schedule_xlsx(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required to read .xlsx schedule files"
        ) from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = None
    for name in wb.sheetnames:
        if re.search(r"进度|schedule|gantt|计划", name, re.I):
            sheet = wb[name]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    fieldnames = [str(c).strip() if c is not None else "" for c in header_row]
    # Prefer first sheet that has a recognizable name + date header
    def _has_schedule_headers(headers: list[str]) -> bool:
        keys = {_header_key(h) for h in headers if h}
        return "name" in keys and ("start" in keys or "end" in keys)

    if not _has_schedule_headers(fieldnames):
        for name in wb.sheetnames:
            cand = wb[name]
            it = cand.iter_rows(values_only=True)
            try:
                hr = next(it)
            except StopIteration:
                continue
            headers = [str(c).strip() if c is not None else "" for c in hr]
            if _has_schedule_headers(headers):
                sheet = cand
                fieldnames = headers
                rows_iter = it
                break
        else:
            raise ValueError(
                f"No schedule sheet with name+date headers in {path}"
            )

    rows: list[dict[str, Any]] = []
    for values in rows_iter:
        row = {
            fieldnames[i]: (values[i] if i < len(values) else None)
            for i in range(len(fieldnames))
            if fieldnames[i]
        }
        rows.append(row)
    return _parse_rows(fieldnames, rows)


def parse_schedule_file(path: Path) -> list[dict[str, str]]:
    path = path.expanduser().resolve()
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return parse_schedule_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_schedule_xlsx(path)
    if suffix == ".mpp":
        raise ValueError(
            "Native .mpp is not read by this Skill. "
            "First run schedule-mpp-excel "
            "(python3 .cursor/skills/schedule-mpp-excel/scripts/convert_schedule.py "
            "--source <file.mpp> --out-dir <00_进度计划>), "
            "then pass the generated *_schedule.csv to --source."
        )
    raise ValueError(f"Unsupported schedule format: {suffix}")


def resolve_schedule_source(project_root: Path, source: Path | None) -> Path:
    if source is not None:
        path = source.expanduser()
        candidates: list[Path] = []
        if path.is_absolute():
            candidates.append(path)
        else:
            candidates.append((project_root / path).resolve())
            candidates.append(Path.cwd() / path)
            candidates.append(path)
        for cand in candidates:
            if cand.is_file():
                return cand.resolve()
        raise FileNotFoundError(
            f"Schedule file not found: {source} "
            f"(tried under project root and cwd)"
        )

    archive = project_root / "项目档案" / "00_进度计划"
    if not archive.is_dir():
        raise FileNotFoundError(
            f"No schedule source and missing directory: {archive}"
        )

    candidates: list[Path] = []
    for pattern in ("*.csv", "*.xlsx", "*.xlsm"):
        candidates.extend(archive.glob(pattern))
    candidates = sorted(
        candidates, key=lambda p: p.stat().st_mtime, reverse=True
    )
    if not candidates:
        raise FileNotFoundError(f"No csv/xlsx under {archive}")
    if len(candidates) == 1:
        return candidates[0].resolve()

    preferred: list[Path] = []
    for pat in ("*进度*", "*schedule*", "*project*", "*Project*"):
        preferred.extend(archive.glob(f"{pat}.csv"))
        preferred.extend(archive.glob(f"{pat}.xlsx"))
    if preferred:
        return sorted(
            preferred, key=lambda p: p.stat().st_mtime, reverse=True
        )[0].resolve()
    return candidates[0].resolve()


def _parse_date(s: str) -> datetime | None:
    s = (s or "").strip()
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        return None


def _add_days(start: str, days: float) -> str:
    dt = _parse_date(start)
    if not dt:
        return start
    return (dt + timedelta(days=int(round(days)))).strftime("%Y-%m-%d")


def _span_days(start: str, end: str) -> int:
    a, b = _parse_date(start), _parse_date(end)
    if not a or not b:
        return 1
    return max(1, (b - a).days)


def _display_name(node_code: str) -> str:
    zh, en = NODE_LABELS.get(node_code, (node_code, node_code))
    return f"{zh} / {en}"


def _match_node_code(task_name: str) -> str | None:
    hay = (task_name or "").lower()
    # Longer / more specific keywords first by sorting codes with more specific set
    best: str | None = None
    best_len = 0
    for code, keywords in NODE_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in hay and len(kw) >= best_len:
                best = code
                best_len = len(kw)
    return best


def _package_tokens(name: str) -> list[str]:
    tokens: list[str] = []
    n = (name or "").strip()
    if not n:
        return tokens
    tokens.append(n.lower())
    for part in re.split(r"[_\-\s/]+", n):
        part = part.strip().lower()
        if len(part) >= 2 and part not in tokens:
            tokens.append(part)
    return tokens


def load_packages_from_tracker(tracker_path: Path) -> list[dict[str, str]]:
    raw = tracker_path.read_text(encoding="utf-8-sig")
    lines = raw.splitlines()
    # Skip "# 最后更新,..." comment row if present
    if lines and lines[0].lstrip().startswith("#"):
        raw = "\n".join(lines[1:])
    reader = csv.DictReader(io.StringIO(raw))
    if not reader.fieldnames:
        return []
    name_key = None
    type_key = None
    for h in reader.fieldnames:
        hn = _normalize_header(h)
        if hn in {"采购包名称", "package_name", "采购包", "package"} and not name_key:
            name_key = h
        if hn in {"采购包类型", "package_type", "类型", "type"} and not type_key:
            type_key = h
    if not name_key:
        # Chinese tracker headers from Skill 1
        for h in reader.fieldnames:
            if "采购包" in h and "名称" in h:
                name_key = h
            if "类型" in h:
                type_key = h
    packages: list[dict[str, str]] = []
    for row in reader:
        if not name_key:
            break
        pname = str(row.get(name_key, "") or "").strip()
        if not pname:
            continue
        ptype = str(row.get(type_key, "") or "").strip() if type_key else "供应"
        type_map = {
            "纯供应": "供应",
            "供应": "供应",
            "供应+安装": "供装",
            "供装": "供装",
            "顾问": "顾问",
            "顾问服务": "顾问",
            "总包": "总包",
        }
        ptype = type_map.get(ptype, ptype or "供应")
        if ptype not in {"供应", "供装", "顾问", "总包"}:
            ptype = "供应"
        packages.append({"package_name": pname, "package_type": ptype})
    return packages


def load_packages_json(text: str) -> list[dict[str, str]]:
    data = json.loads(text)
    if not isinstance(data, list):
        raise ValueError("--packages-json must be a JSON array")
    out: list[dict[str, str]] = []
    for item in data:
        if not isinstance(item, dict):
            continue
        name = str(
            item.get("package_name") or item.get("name") or ""
        ).strip()
        if not name:
            continue
        ptype = str(item.get("package_type") or item.get("type") or "供应").strip()
        type_map = {
            "纯供应": "供应",
            "供应+安装": "供装",
            "顾问服务": "顾问",
        }
        ptype = type_map.get(ptype, ptype)
        if ptype not in {"供应", "供装", "顾问", "总包"}:
            ptype = "供应"
        out.append({"package_name": name, "package_type": ptype})
    return out


def find_tracker(project_root: Path) -> Path | None:
    matches = list(project_root.glob("*_procurement_packages_status_tracker.csv"))
    if matches:
        return sorted(matches, key=lambda p: p.stat().st_mtime, reverse=True)[0]
    legacy = project_root / "tracker.csv"
    return legacy if legacy.is_file() else None


def assign_package(
    task: dict[str, str],
    packages: list[dict[str, str]],
) -> tuple[str | None, str | None]:
    """Return (package_name, package_type) or (_project / None)."""
    hint = (task.get("package_hint") or "").strip()
    raw = (task.get("raw_name") or task.get("name") or "").strip()
    full = (task.get("name") or "").strip()
    node = _match_node_code(raw) or _match_node_code(full)

    if hint:
        for pkg in packages:
            if pkg["package_name"].lower() == hint.lower():
                return pkg["package_name"], pkg["package_type"]
        # hint present but not in list — treat as package name, default type
        return hint, "供应"

    if node and node.startswith("project_"):
        return "_project", "_project"

    for pkg in packages:
        tokens = _package_tokens(pkg["package_name"])
        hay = full.lower()
        if any(t in hay for t in tokens):
            return pkg["package_name"], pkg["package_type"]

    if node and node.startswith("project_"):
        return "_project", "_project"
    if node:
        # keyword hit but no package — leave unmapped unless project-like
        return None, None
    return None, None


def _make_node(
    *,
    node_code: str,
    window_start: str,
    window_end: str,
    source: str,
    confidence: float,
    basis: str,
    package_name: str,
    package_type: str,
    risks: list[str] | None = None,
    needs_confirm: bool = False,
) -> dict[str, Any]:
    zh, en = NODE_LABELS.get(node_code, (node_code, node_code))
    node: dict[str, Any] = {
        "node_code": node_code,
        "name": _display_name(node_code),
        "name_zh": zh,
        "name_en": en,
        "window_start": window_start,
        "window_end": window_end or window_start,
        "source": source,
        "confidence": round(confidence, 2),
        "basis": basis,
        "package_name": package_name,
        "package_type": package_type,
        "risks": risks or [],
    }
    if needs_confirm:
        node["needs_confirm"] = True
    return node


def analyze(
    milestones: list[dict[str, str]],
    packages: list[dict[str, str]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    pending: list[dict[str, Any]] = []
    critical: list[dict[str, Any]] = []

    # Discover packages from schedule if empty
    if not packages:
        discovered: dict[str, str] = {}
        for task in milestones:
            hint = (task.get("package_hint") or "").strip()
            if hint:
                discovered.setdefault(hint, "供应")
            else:
                # try extract first token before space as weak package — skip
                pass
        packages = [
            {"package_name": k, "package_type": v} for k, v in discovered.items()
        ]
        if packages:
            pending.append(
                {
                    "code": "package_type_defaulted",
                    "message": "采购包类型未提供，已默认「供应」，请确认供装/顾问/总包。",
                    "context": {"packages": [p["package_name"] for p in packages]},
                }
            )

    # Ensure all packages listed
    pkg_by_name = {p["package_name"]: p for p in packages}

    # Group tasks
    project_tasks: list[dict[str, str]] = []
    by_package: dict[str, list[dict[str, str]]] = {p["package_name"]: [] for p in packages}
    coarse_bars: dict[str, dict[str, str]] = {}

    for task in milestones:
        pname, ptype = assign_package(task, packages)
        node = _match_node_code(task.get("raw_name") or "") or _match_node_code(
            task.get("name") or ""
        )

        if pname == "_project" or (node and node.startswith("project_")):
            project_tasks.append(task)
            continue

        if pname is None:
            if node:
                pending.append(
                    {
                        "code": "unmapped_task",
                        "message": f"任务「{task.get('name')}」命中节点语义但无法归属采购包，请指定包名。",
                        "context": {"task": task.get("name"), "node_code": node},
                    }
                )
            else:
                pending.append(
                    {
                        "code": "unmapped_task",
                        "message": f"任务「{task.get('name')}」无法归属采购包。",
                        "context": {"task": task.get("name")},
                    }
                )
            continue

        if pname not in by_package:
            by_package[pname] = []
            pkg_by_name.setdefault(
                pname, {"package_name": pname, "package_type": ptype or "供应"}
            )
        by_package[pname].append(task)

        if not node:
            # coarse bar candidate
            existing = coarse_bars.get(pname)
            if not existing:
                coarse_bars[pname] = task
            else:
                # widen window
                s = min(existing["date"], task["date"])
                e = max(
                    existing.get("end_date") or existing["date"],
                    task.get("end_date") or task["date"],
                )
                coarse_bars[pname] = {
                    **task,
                    "date": s,
                    "end_date": e,
                    "name": existing.get("name") or task.get("name") or "",
                }

    # Project-level nodes (dedupe by node_code)
    project_anchors: dict[str, dict[str, str]] = {}
    seen_project: set[str] = set()
    for task in milestones:
        node = _match_node_code(task.get("raw_name") or "") or _match_node_code(
            task.get("name") or ""
        )
        if not (node and node.startswith("project_")):
            continue
        if node in seen_project:
            continue
        seen_project.add(node)
        project_anchors[node] = task
        critical.append(
            _make_node(
                node_code=node,
                window_start=task["date"],
                window_end=task.get("end_date") or task["date"],
                source="explicit",
                confidence=0.95,
                basis=f"Matched task 「{task.get('name')}」",
                package_name="_project",
                package_type="_project",
            )
        )

    for code in DEFAULT_NODES["_project"]:
        if code in project_anchors:
            continue
        pending.append(
            {
                "code": "missing_anchor",
                "message": f"项目级节点 {code}（{NODE_LABELS.get(code, (code,))[0]}）未在进度计划中找到。",
                "context": {"node_code": code},
            }
        )

    handover = project_anchors.get("project_handover") or project_anchors.get(
        "project_completion"
    )
    handover_date = (
        (handover.get("end_date") or handover["date"]) if handover else ""
    )

    # Per-package expand
    for pname, tasks in by_package.items():
        pkg = pkg_by_name.get(pname) or {
            "package_name": pname,
            "package_type": "供应",
        }
        ptype = pkg["package_type"]
        if ptype == "总包":
            for task in tasks:
                node = _match_node_code(task.get("raw_name") or "") or _match_node_code(
                    task.get("name") or ""
                )
                if node and node.startswith("project_"):
                    continue
                critical.append(
                    _make_node(
                        node_code=node or "gc_milestone",
                        window_start=task["date"],
                        window_end=task.get("end_date") or task["date"],
                        source="explicit",
                        confidence=0.9 if node else 0.7,
                        basis=f"GC schedule row 「{task.get('name')}」",
                        package_name=pname,
                        package_type=ptype,
                    )
                )
            continue

        explicit_by_code: dict[str, dict[str, str]] = {}
        for task in tasks:
            node = _match_node_code(task.get("raw_name") or "") or _match_node_code(
                task.get("name") or ""
            )
            if node and not node.startswith("project_"):
                explicit_by_code[node] = task

        for code, task in explicit_by_code.items():
            conf = 0.95
            raw = (task.get("raw_name") or "").lower()
            if pname.lower() not in (task.get("name") or "").lower() and not task.get(
                "package_hint"
            ):
                conf = 0.75
            critical.append(
                _make_node(
                    node_code=code,
                    window_start=task["date"],
                    window_end=task.get("end_date") or task["date"],
                    source="explicit",
                    confidence=conf,
                    basis=f"Matched task 「{task.get('name')}」 keywords → {code}",
                    package_name=pname,
                    package_type=ptype,
                )
            )

        required = list(DEFAULT_NODES.get(ptype, ["delivery"]))
        bar = coarse_bars.get(pname)
        # If no coarse bar but has explicit nodes, use union of explicit as bar for missing
        if not bar and explicit_by_code:
            dates = []
            for t in explicit_by_code.values():
                dates.append(t["date"])
                dates.append(t.get("end_date") or t["date"])
            bar = {
                "name": f"{pname} (from explicit nodes)",
                "date": min(dates),
                "end_date": max(dates),
            }

        fractions = INFER_FRACTIONS.get(ptype, {})

        for code in required:
            if code in explicit_by_code:
                continue

            # Prefer placing missing early nodes before the next known explicit successor
            successor = None
            try:
                idx = required.index(code)
            except ValueError:
                idx = -1
            if idx >= 0:
                for later in required[idx + 1 :]:
                    if later in explicit_by_code:
                        successor = explicit_by_code[later]
                        break

            if successor and code in {"delivery", "site_access", "po_confirm", "fat", "shipment"}:
                end = successor["date"]
                start_anchor = project_anchors.get("project_start")
                start = start_anchor["date"] if start_anchor else end
                # Keep a short window ending at successor start
                dur = _span_days(start, end)
                lead = max(7, min(21, dur // 4 or 7))
                ws = _add_days(end, -lead)
                if start_anchor and ws < start_anchor["date"]:
                    ws = start_anchor["date"]
                we = end
                conf = 0.72
                basis = (
                    f"Inferred {code} ending at successor "
                    f"{successor.get('name')} start {end}"
                )
                needs = False
            elif bar:
                s, e = bar["date"], bar.get("end_date") or bar["date"]
                dur = _span_days(s, e)
                frac = fractions.get(code, (0.0, 1.0))
                ws = _add_days(s, frac[0] * dur)
                we = _add_days(s, frac[1] * dur)
                if handover_date and code in {
                    "final_accept",
                    "deliverable_accept",
                    "prelim_accept",
                }:
                    if we > handover_date:
                        we = handover_date
                conf = 0.7
                basis = (
                    f"Inferred from coarse bar 「{bar.get('name')}」 "
                    f"[{s} ~ {e}] using {ptype} template fraction {frac}"
                )
                needs = False
            elif handover_date or project_anchors.get("project_start"):
                start_a = project_anchors.get("project_start")
                s = start_a["date"] if start_a else ""
                e = handover_date or s
                if not s or not e:
                    pending.append(
                        {
                            "code": "low_confidence",
                            "message": f"包 {pname} 缺少 {code}，且无粗条/锚点可推断。",
                            "context": {"package_name": pname, "node_code": code},
                        }
                    )
                    continue
                dur = _span_days(s, e)
                frac = fractions.get(code, (0.4, 0.6))
                ws = _add_days(s, frac[0] * dur)
                we = _add_days(s, frac[1] * dur)
                conf = 0.48
                basis = (
                    f"Inferred from project anchors only "
                    f"(start={s}, handover={e}); confirm with PM"
                )
                needs = True
            else:
                pending.append(
                    {
                        "code": "missing_anchor",
                        "message": f"包 {pname} 缺少节点 {code}，且无粗条或项目锚点。",
                        "context": {"package_name": pname, "node_code": code},
                    }
                )
                continue

            node = _make_node(
                node_code=code,
                window_start=ws,
                window_end=we,
                source="inferred",
                confidence=conf,
                basis=basis,
                package_name=pname,
                package_type=ptype,
                needs_confirm=needs,
            )
            if conf < 0.5:
                pending.append(
                    {
                        "code": "low_confidence",
                        "message": f"包 {pname} 节点 {code} 置信度 {conf}，请确认窗口 {ws} ~ {we}。",
                        "context": {
                            "package_name": pname,
                            "node_code": code,
                            "window_start": ws,
                            "window_end": we,
                        },
                    }
                )
            critical.append(node)

    # Align site_access to delivery end when both exist and site was inferred
    by_pkg_nodes: dict[str, list[dict[str, Any]]] = {}
    for n in critical:
        by_pkg_nodes.setdefault(n["package_name"], []).append(n)

    for pname, nodes in by_pkg_nodes.items():
        delivery = next((n for n in nodes if n["node_code"] == "delivery"), None)
        site = next((n for n in nodes if n["node_code"] == "site_access"), None)
        if (
            delivery
            and site
            and site["source"] == "inferred"
            and delivery["source"] == "explicit"
        ):
            site["window_start"] = delivery["window_end"]
            site["window_end"] = delivery["window_end"]
            site["source"] = "aligned"
            site["confidence"] = 0.8
            site["basis"] = (
                f"Aligned site_access to delivery end {delivery['window_end']}"
            )

    # Conflict checks
    for pname, nodes in by_pkg_nodes.items():
        if pname == "_project":
            continue
        ordered = sorted(
            nodes, key=lambda n: NODE_ORDER.get(n["node_code"], 999)
        )
        for i in range(1, len(ordered)):
            prev, cur = ordered[i - 1], ordered[i]
            if cur["window_start"] < prev["window_start"]:
                msg = (
                    f"顺序冲突：{cur['node_code']} 开始 {cur['window_start']} "
                    f"早于 {prev['node_code']} 开始 {prev['window_start']}"
                )
                cur.setdefault("risks", []).append(msg)
        delivery = next((n for n in nodes if n["node_code"] == "delivery"), None)
        install = next(
            (n for n in nodes if n["node_code"] == "install_complete"), None
        )
        if (
            delivery
            and install
            and delivery["window_end"] > install["window_end"]
        ):
            msg = (
                f"到货窗口结束晚于安装完成结束："
                f"delivery {delivery['window_end']} > install_complete {install['window_end']}"
            )
            delivery.setdefault("risks", []).append(msg)
        if handover_date:
            for n in nodes:
                if n["node_code"] in {
                    "final_accept",
                    "prelim_accept",
                    "deliverable_accept",
                } and n["window_end"] > handover_date:
                    msg = (
                        f"验收结束晚于项目移交/竣工："
                        f"{n['node_code']} {n['window_end']} > {handover_date}"
                    )
                    n.setdefault("risks", []).append(msg)

    # Sort output
    critical.sort(
        key=lambda n: (
            n["package_name"] != "_project",
            n["package_name"],
            NODE_ORDER.get(n["node_code"], 999),
            n["window_start"],
        )
    )
    return critical, pending


def milestones_for_skill1(raw: list[dict[str, str]]) -> list[dict[str, str]]:
    """Strip helper fields for Skill 1-compatible milestones[]."""
    out: list[dict[str, str]] = []
    for item in raw:
        row: dict[str, str] = {
            "name": item["name"],
            "date": item["date"],
        }
        end = item.get("end_date") or ""
        if end and end != item["date"]:
            row["end_date"] = end
        out.append(row)
    return out


def _decade_label(iso_date: str) -> str:
    """Map YYYY-MM-DD → 'YYYY年M月上/中/下旬'."""
    dt = _parse_date(iso_date)
    if not dt:
        return iso_date or ""
    if dt.day <= 10:
        decade = "上旬"
    elif dt.day <= 20:
        decade = "中旬"
    else:
        decade = "下旬"
    return f"{dt.year}年{dt.month}月{decade}"


def render_markdown(
    *,
    analyzed_at: str,
    source_file: str,
    critical: list[dict[str, Any]],
    pending: list[dict[str, Any]],
) -> str:
    """Human summary: payment ordinal + completion decade (presentation.md)."""
    lines = [
        "# 关键节点摘要",
        f"分析时间: {analyzed_at}",
        f"源文件: {source_file}",
        "",
        "说明: 下列时间为**完成时点**（请款对齐 `window_end`）；旬粒度便于人读。",
        "",
    ]
    by_pkg: dict[str, list[dict[str, Any]]] = {}
    for n in critical:
        by_pkg.setdefault(n["package_name"], []).append(n)

    def _emit_pkg(title: str, nodes: list[dict[str, Any]]) -> None:
        lines.append(f"## {title}")
        for i, n in enumerate(nodes, start=1):
            end = n.get("window_end") or n.get("window_start") or ""
            decade = _decade_label(str(end))
            flag = " ⚠ confirm" if n.get("needs_confirm") else ""
            inferred = "（倒推）" if n.get("source") == "inferred" else ""
            lines.append(
                f"{i}. {n['name']} · 完成 {end}（{decade}）{inferred} | "
                f"{n['source']} | conf={n['confidence']}{flag} | {n['basis']}"
            )
            for r in n.get("risks") or []:
                lines.append(f"   - risk: {r}")
        lines.append("")

    if "_project" in by_pkg:
        _emit_pkg("项目级", by_pkg["_project"])

    for pname in sorted(k for k in by_pkg if k != "_project"):
        nodes = by_pkg[pname]
        ptype = nodes[0].get("package_type", "")
        _emit_pkg(f"{pname}（{ptype}）", nodes)

    if pending:
        lines.append("## Pending")
        for p in pending:
            lines.append(f"- [{p.get('code')}] {p.get('message')}")
        lines.append("")

    return "\n".join(lines)


def load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def save_json(path: Path, data: dict[str, Any]) -> None:
    if path.name == "pending.json":
        raise PermissionError(
            "进度的待确认必须写在 schedule.json 里，不能写进 pending.json。"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def _render_timeline_files(
    root: Path,
    *,
    package_name: str | None = None,
    package_dir: Path | None = None,
) -> dict[str, Any]:
    """Generate project (+ optional per-package) timeline HTML via render_timeline.py."""
    from render_timeline import write_timeline

    project = write_timeline(root)
    out: dict[str, Any] = {"project": project}
    if package_name and package_dir:
        out["package"] = write_timeline(
            root,
            package_filter=package_name,
            package_dir=package_dir,
            highlight_package=package_name,
        )
    return out


def run(
    project_root: Path,
    source: Path | None,
    packages: list[dict[str, str]],
    *,
    write_md: bool = False,
    write_timeline: bool = False,
    timeline_package: str | None = None,
    timeline_package_dir: Path | None = None,
) -> dict[str, Any]:
    root = project_root.expanduser().resolve()
    src = resolve_schedule_source(root, source)
    milestones_raw = parse_schedule_file(src)
    if not milestones_raw:
        raise ValueError(f"No milestones parsed from {src}")

    critical, pending = analyze(milestones_raw, packages)

    skill_dir = root / ".skill"
    schedule_path = skill_dir / "schedule.json"
    existing = load_json(schedule_path)

    if src.is_relative_to(root):
        rel = str(src.relative_to(root))
    else:
        rel = str(src)

    doc: dict[str, Any] = dict(existing) if isinstance(existing, dict) else {}
    doc["milestones"] = milestones_for_skill1(milestones_raw)
    doc["critical_nodes"] = critical
    doc["pending"] = pending
    doc["packages"] = packages
    doc["source_file"] = rel
    doc["analyzed_at"] = now_iso()
    doc["note"] = "critical_nodes from schedule-critical-nodes Skill"
    save_json(schedule_path, doc)

    md_path = None
    if write_md:
        md_path = skill_dir / "critical_nodes.md"
        md_path.write_text(
            render_markdown(
                analyzed_at=doc["analyzed_at"],
                source_file=rel,
                critical=critical,
                pending=pending,
            ),
            encoding="utf-8",
        )

    timeline_info = None
    if write_timeline:
        timeline_info = _render_timeline_files(
            root,
            package_name=timeline_package,
            package_dir=timeline_package_dir,
        )

    return {
        "project_root": str(root),
        "source_file": str(src),
        "schedule_json": str(schedule_path),
        "critical_nodes_md": str(md_path) if md_path else None,
        "timeline": timeline_info,
        "milestone_count": len(milestones_raw),
        "critical_node_count": len(critical),
        "pending_count": len(pending),
        "inferred_count": sum(1 for n in critical if n["source"] == "inferred"),
        "explicit_count": sum(1 for n in critical if n["source"] == "explicit"),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Analyze schedule → .skill/schedule.json critical_nodes"
    )
    parser.add_argument("--project-root", required=True)
    parser.add_argument("--source", default=None, help="csv/xlsx path")
    parser.add_argument("--packages-json", default=None)
    parser.add_argument("--tracker", default=None, help="Tracker CSV path")
    parser.add_argument(
        "--write-md",
        action="store_true",
        help="Also write .skill/critical_nodes.md",
    )
    parser.add_argument(
        "--write-timeline",
        action="store_true",
        help="Also write 关键节点时间线.html (presentation.md §3)",
    )
    parser.add_argument(
        "--timeline-package",
        default=None,
        help="With --write-timeline: also write filtered timeline into --timeline-package-dir",
    )
    parser.add_argument(
        "--timeline-package-dir",
        default=None,
        help="Package folder for filtered timeline HTML",
    )
    args = parser.parse_args(argv)

    root = Path(args.project_root).expanduser().resolve()
    packages: list[dict[str, str]] = []

    if args.packages_json:
        packages = load_packages_json(args.packages_json)
    else:
        tracker = (
            Path(args.tracker).expanduser().resolve()
            if args.tracker
            else find_tracker(root)
        )
        if tracker and tracker.is_file():
            packages = load_packages_from_tracker(tracker)

    source = Path(args.source) if args.source else None
    pkg_dir = (
        Path(args.timeline_package_dir).expanduser().resolve()
        if args.timeline_package_dir
        else None
    )
    try:
        result = run(
            root,
            source,
            packages,
            write_md=args.write_md,
            write_timeline=args.write_timeline,
            timeline_package=args.timeline_package,
            timeline_package_dir=pkg_dir,
        )
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
