#!/usr/bin/env python3
"""Machine-readable CSV / flat xlsx export."""

from __future__ import annotations

import csv
from pathlib import Path

from readers import TaskRow

# First four columns: 任务名称, 开始日期, 结束日期, 采购包 (stable machine-readable contract)
CSV_HEADERS = [
    "任务名称",
    "开始日期",
    "结束日期",
    "采购包",
    "ID",
    "WBS",
    "大纲级别",
    "工期",
    "前置任务",
    "完成百分比",
    "资源名称",
    "是否摘要",
    "是否里程碑",
    "是否关键",
]


def write_structured_csv(tasks: list[TaskRow], out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_HEADERS)
        w.writeheader()
        for t in tasks:
            w.writerow(_row_dict(t))
    return out_path


def write_structured_xlsx(tasks: list[TaskRow], out_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    for col, h in enumerate(CSV_HEADERS, start=1):
        cell = ws.cell(1, col, h)
        cell.font = Font(bold=True)
    for r_i, t in enumerate(tasks, start=2):
        d = _row_dict(t)
        for col, h in enumerate(CSV_HEADERS, start=1):
            ws.cell(r_i, col, d[h])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _row_dict(t: TaskRow) -> dict:
    return {
        "任务名称": t.name,
        "开始日期": t.start.isoformat() if t.start else "",
        "结束日期": t.finish.isoformat() if t.finish else "",
        "采购包": t.package,
        "ID": t.task_id,
        "WBS": t.wbs,
        "大纲级别": t.outline_level,
        "工期": t.duration,
        "前置任务": t.predecessors,
        "完成百分比": int(t.percent_complete) if t.percent_complete else 0,
        "资源名称": t.resources,
        "是否摘要": "是" if t.is_summary else "否",
        "是否里程碑": "是" if t.is_milestone else "否",
        "是否关键": "是" if t.is_critical else "否",
    }
