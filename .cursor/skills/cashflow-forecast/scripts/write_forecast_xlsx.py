#!/usr/bin/env python3
"""Write the Skill 4 cashflow workbook: one row per package, months across."""

from __future__ import annotations

from pathlib import Path
from typing import Any

INFO_HEADERS = [
    "序号",
    "采购包",
    "供应商",
    "合同金额（不含税）",
    "已确认变更",
    "PO号",
    "请款条件",
    "待确认金额",
]

MONEY_HEADERS = {
    "合同金额（不含税）",
    "已确认变更",
    "待确认金额",
}

COL_WIDTHS = {
    "序号": 6,
    "采购包": 14,
    "供应商": 24,
    "合同金额（不含税）": 16,
    "已确认变更": 14,
    "PO号": 16,
    "请款条件": 38,
    "待确认金额": 14,
    "应请": 12,
    "实请": 12,
}


def _theme() -> dict[str, str]:
    return {
        "ink": "0F172A",
        "ink_muted": "64748B",
        "header_bg": "334155",
        "header_text": "FFFFFF",
        "group_bg": "F1F5F9",
        "group_text": "475569",
        "zebra": "F8FAFC",
        "actual_hdr": "1D4ED8",
        "actual_hdr_bg": "DBEAFE",
        "total_bg": "E2E8F0",
        "border": "E2E8F0",
        "font": "Microsoft YaHei",
    }


def _month_label(month_key: str) -> str:
    y, m = month_key.split("-")
    return f"{int(y)}年{int(m)}月"


def write_forecast_xlsx(
    path: Path,
    doc: dict[str, Any],
    project_name: str,
    check_command: Path | None = None,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    theme = _theme()
    thin = Side(style="thin", color=theme["border"])
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731

    fonts = {
        "meta": Font(name=theme["font"], size=10, color=theme["ink_muted"]),
        "title": Font(name=theme["font"], size=12, bold=True, color=theme["ink"]),
        "group": Font(name=theme["font"], size=10, bold=True, color=theme["group_text"]),
        "header": Font(name=theme["font"], size=10, bold=True, color=theme["header_text"]),
        "expected": Font(name=theme["font"], size=9, bold=True, color=theme["header_text"]),
        "actual": Font(name=theme["font"], size=9, bold=True, color=theme["actual_hdr"]),
        "body": Font(name=theme["font"], size=10, color=theme["ink"]),
        "button": Font(name=theme["font"], size=10, bold=True, color=theme["header_text"]),
        "link": Font(name=theme["font"], size=10, color=theme["actual_hdr"], underline="single"),
        "total": Font(name=theme["font"], size=10, bold=True, color=theme["ink"]),
    }
    fills = {
        "header": fill(theme["header_bg"]),
        "group": fill(theme["group_bg"]),
        "zebra": fill(theme["zebra"]),
        "expected": fill(theme["header_bg"]),
        "actual": fill(theme["actual_hdr_bg"]),
        "button": fill(theme["actual_hdr"]),
        "total": fill(theme["total_bg"]),
    }
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    months: list[str] = list(doc.get("calendar", {}).get("months") or [])
    packages: list[dict[str, Any]] = list(doc.get("packages") or [])
    info_n = len(INFO_HEADERS)
    last_col = info_n + max(1, len(months) * 2)

    wb = Workbook()
    ws = wb.active
    ws.title = "现金流"

    as_of = doc.get("as_of") or ""
    lag = doc.get("payment_lag_days") or 30
    ws.cell(
        1,
        1,
        f"现金流预测 · {project_name} · 截至 {as_of} · 不含税 · 出账滞后 {lag} 天",
    )
    ws.cell(1, 1).font = fonts["title"]
    title_end = min(7, max(3, last_col - 2))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=title_end)
    btn_start = title_end + 1
    btn_end = min(btn_start + 1, last_col) if last_col > btn_start else btn_start
    btn = ws.cell(1, btn_start, "检查台账更新")
    btn.font = fonts["button"]
    btn.fill = fills["button"]
    btn.alignment = center
    btn.hyperlink = "#台账来源!A1"
    if btn_end > btn_start:
        ws.merge_cells(start_row=1, start_column=btn_start, end_row=1, end_column=btn_end)
        extra = ws.cell(1, btn_end)
        extra.fill = fills["button"]
        extra.font = fonts["button"]
    ws.row_dimensions[1].height = 26

    for col in range(1, info_n + 1):
        cell = ws.cell(2, col, "基础信息" if col == 1 else None)
        cell.fill = fills["group"]
        cell.font = fonts["group"]
        cell.alignment = center
        cell.border = border
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=info_n)

    for i, month in enumerate(months):
        start = info_n + 1 + i * 2
        cell = ws.cell(2, start, _month_label(month))
        cell.fill = fills["group"]
        cell.font = fonts["group"]
        cell.alignment = center
        cell.border = border
        ws.cell(2, start + 1).fill = fills["group"]
        ws.cell(2, start + 1).border = border
        ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=start + 1)
    ws.row_dimensions[2].height = 22

    for col, header in enumerate(INFO_HEADERS, start=1):
        cell = ws.cell(3, col, header)
        cell.font = fonts["header"]
        cell.fill = fills["header"]
        cell.alignment = center
        cell.border = border
    for i, _month in enumerate(months):
        start = info_n + 1 + i * 2
        exp = ws.cell(3, start, "应请")
        exp.font = fonts["expected"]
        exp.fill = fills["expected"]
        exp.alignment = center
        exp.border = border
        act = ws.cell(3, start + 1, "实请")
        act.font = fonts["actual"]
        act.fill = fills["actual"]
        act.alignment = center
        act.border = border
    ws.row_dimensions[3].height = 28

    data_start = 4

    def _write_money(cell: Any, value: Any) -> None:
        if value in (None, ""):
            cell.value = None
            return
        try:
            num = float(value)
        except (TypeError, ValueError):
            cell.value = value
            return
        if num == 0:
            cell.value = None
            return
        cell.value = num
        cell.number_format = "#,##0.00"

    for row_idx, pkg in enumerate(packages, start=data_start):
        zebra = (row_idx - data_start) % 2 == 1
        row_fill = fills["zebra"] if zebra else None
        values = [
            pkg.get("seq") or "",
            pkg.get("package") or "",
            pkg.get("supplier") or "",
            pkg.get("amount_excl_tax"),
            pkg.get("changes_excl_tax"),
            pkg.get("po_number") or "",
            pkg.get("payment_terms") or "",
            pkg.get("pending_excl_tax"),
        ]
        for col, (header, value) in enumerate(zip(INFO_HEADERS, values), start=1):
            cell = ws.cell(row_idx, col)
            if header in MONEY_HEADERS:
                _write_money(cell, value)
                cell.alignment = right
            else:
                cell.value = value if value not in ("", None) else None
                cell.alignment = center if header == "序号" else left
            cell.font = fonts["body"]
            cell.border = border
            if row_fill:
                cell.fill = row_fill

        month_map = pkg.get("months") or {}
        for i, month in enumerate(months):
            bucket = month_map.get(month) or {}
            exp_cell = ws.cell(row_idx, info_n + 1 + i * 2)
            act_cell = ws.cell(row_idx, info_n + 2 + i * 2)
            _write_money(exp_cell, bucket.get("expected"))
            _write_money(act_cell, bucket.get("actual"))
            exp_cell.alignment = right
            act_cell.alignment = right
            exp_cell.font = fonts["body"]
            act_cell.font = fonts["body"]
            exp_cell.border = border
            act_cell.border = border
            if row_fill:
                exp_cell.fill = row_fill
                act_cell.fill = row_fill
        ws.row_dimensions[row_idx].height = 22

    total_row = data_start + len(packages)
    last_data = max(data_start, total_row - 1)
    for col, header in enumerate(INFO_HEADERS, start=1):
        cell = ws.cell(total_row, col)
        cell.font = fonts["total"]
        cell.fill = fills["total"]
        cell.border = border
        if col == 1:
            cell.value = "总计"
            cell.alignment = center
        elif header in MONEY_HEADERS and packages:
            letter = get_column_letter(col)
            cell.value = f"=SUM({letter}{data_start}:{letter}{last_data})"
            cell.number_format = "#,##0.00"
            cell.alignment = right
        else:
            cell.alignment = left
    for i, _month in enumerate(months):
        for offset in (0, 1):
            col = info_n + 1 + i * 2 + offset
            cell = ws.cell(total_row, col)
            cell.font = fonts["total"]
            cell.fill = fills["total"]
            cell.border = border
            cell.alignment = right
            if packages:
                letter = get_column_letter(col)
                cell.value = f"=SUM({letter}{data_start}:{letter}{last_data})"
                cell.number_format = "#,##0.00"
    ws.row_dimensions[total_row].height = 24

    for col, header in enumerate(INFO_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = COL_WIDTHS.get(header, 12)
    for i, _month in enumerate(months):
        ws.column_dimensions[get_column_letter(info_n + 1 + i * 2)].width = COL_WIDTHS["应请"]
        ws.column_dimensions[get_column_letter(info_n + 2 + i * 2)].width = COL_WIDTHS["实请"]

    ws.freeze_panes = ws.cell(data_start, info_n + 1)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = theme["header_bg"]

    _write_detail_sheet(wb, doc, fonts, fills, border, left, right, center)
    _write_sources_sheet(wb, doc, fonts, fills, border, left, center, check_command)
    warnings = doc.get("warnings") or []
    pending = doc.get("pending") or []
    if warnings or pending:
        _write_notes_sheet(wb, warnings, pending, fonts, left)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _write_detail_sheet(
    wb: Any,
    doc: dict[str, Any],
    fonts: dict,
    fills: dict,
    border: Any,
    left: Any,
    right: Any,
    center: Any,
) -> None:
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("节点明细")
    headers = [
        "采购包",
        "节点",
        "比例",
        "应请金额",
        "计划出账月",
        "实请金额",
        "实际请款月",
        "状态",
        "日期来源",
    ]
    ws.cell(1, 1, "按付款节点拆开的明细（主表在「现金流」；本表供核对）")
    ws.cell(1, 1).font = fonts["meta"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(2, col, header)
        cell.font = fonts["header"]
        cell.fill = fills["header"]
        cell.alignment = center
        cell.border = border
    for row_idx, item in enumerate(doc.get("line_items") or [], start=3):
        values = [
            item.get("package"),
            item.get("node"),
            item.get("pct"),
            item.get("expected_amount_excl_tax"),
            item.get("planned_cash_out_month"),
            item.get("actual_amount_excl_tax"),
            item.get("actual_month"),
            item.get("status"),
            item.get("date_source"),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col, value if value not in ("", None) else None)
            cell.font = fonts["body"]
            cell.border = border
            if col in (4, 6) and isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = right
            elif col in (5, 7, 8):
                cell.alignment = center
            else:
                cell.alignment = left
    widths = [14, 28, 8, 14, 14, 14, 14, 12, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


def _file_uri(path_str: str) -> str:
    p = Path(str(path_str))
    return p.expanduser().resolve().as_uri()


def _write_sources_sheet(
    wb: Any,
    doc: dict[str, Any],
    fonts: dict,
    fills: dict,
    border: Any,
    left: Any,
    center: Any,
    check_command: Path | None,
) -> None:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("台账来源", 1)
    ws.cell(1, 1, "这张现金流表参考了哪些台账").font = fonts["title"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.cell(
        2,
        1,
        "点下面蓝色按钮：核对生成之后这些台账有没有人改过。有更新会重算 Excel。"
        "请先关掉本表再点，跑完重新打开。Excel 若拦截链接，去项目文件夹双击「检查现金流台账.command」。",
    ).font = fonts["meta"]
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 48

    btn = ws.cell(4, 1, "一键检查是否有更新")
    btn.font = fonts["button"]
    btn.fill = fills["button"]
    btn.alignment = center
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=2)
    ws.cell(4, 2).fill = fills["button"]
    if check_command and check_command.is_file():
        btn.hyperlink = _file_uri(str(check_command))
    ws.row_dimensions[4].height = 28

    headers = ["台账", "打开文件", "本表生成时的修改时间", "状态"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(6, col, header)
        cell.font = fonts["header"]
        cell.fill = fills["header"]
        cell.alignment = center
        cell.border = border

    sources = list(doc.get("sources") or [])
    for row_idx, src in enumerate(sources, start=7):
        ws.cell(row_idx, 1, src.get("label") or src.get("rel") or "").font = fonts["body"]
        ws.cell(row_idx, 1).alignment = left
        ws.cell(row_idx, 1).border = border
        open_cell = ws.cell(row_idx, 2, src.get("rel") or Path(str(src.get("path") or "")).name)
        open_cell.border = border
        open_cell.alignment = left
        path_str = str(src.get("path") or "")
        if path_str and not src.get("missing") and Path(path_str).is_file():
            open_cell.hyperlink = _file_uri(path_str)
            open_cell.font = fonts["link"]
        else:
            open_cell.font = fonts["body"]
            if src.get("missing"):
                open_cell.value = "（当时还没有这个文件）"
        ws.cell(row_idx, 3, src.get("mtime") or "—").font = fonts["body"]
        ws.cell(row_idx, 3).alignment = center
        ws.cell(row_idx, 3).border = border
        status = "未找到" if src.get("missing") else "已纳入计算"
        ws.cell(row_idx, 4, status).font = fonts["body"]
        ws.cell(row_idx, 4).alignment = center
        ws.cell(row_idx, 4).border = border

    widths = [28, 42, 24, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A7"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1D4ED8"


def _write_notes_sheet(
    wb: Any,
    warnings: list[dict[str, Any]],
    pending: list[dict[str, Any]],
    fonts: dict,
    left: Any,
) -> None:
    ws = wb.create_sheet("待确认")
    ws.cell(1, 1, "待你拍板 / 警告").font = fonts["title"]
    row = 3
    if warnings:
        ws.cell(row, 1, "警告").font = fonts["group"]
        row += 1
        for w in warnings:
            ws.cell(row, 1, w.get("message") or w.get("code") or str(w)).font = fonts["body"]
            ws.cell(row, 1).alignment = left
            row += 1
        row += 1
    if pending:
        ws.cell(row, 1, "待确认").font = fonts["group"]
        row += 1
        for p in pending:
            ws.cell(row, 1, p.get("message") or p.get("reason") or str(p)).font = fonts["body"]
            ws.cell(row, 1).alignment = left
            row += 1
    ws.column_dimensions["A"].width = 80
