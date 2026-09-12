#!/usr/bin/env python3
"""Smoke test Skill 4 against its own sample project."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

_SKILL_ROOT = Path(__file__).resolve().parents[1]
_BUILD = _SKILL_ROOT / "scripts" / "build_forecast.py"
_SAMPLE = _SKILL_ROOT / "samples" / "demo_project"
if str(_SKILL_ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(_SKILL_ROOT / "scripts"))
_FOUNDATION = _SKILL_ROOT.parents[1] / "lib" / "project-foundation" / "scripts"
if str(_FOUNDATION) not in sys.path:
    sys.path.insert(0, str(_FOUNDATION))


def main() -> int:
    if not _SAMPLE.is_dir():
        raise SystemExit(f"Missing sample: {_SAMPLE}")
    proc = subprocess.run(
        [
            sys.executable,
            str(_BUILD),
            "--project-root",
            str(_SAMPLE),
            "--as-of",
            "2026-09-06",
            "--write",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if proc.returncode != 0:
        raise SystemExit(
            f"build_forecast failed for {_SAMPLE}:\n{proc.stderr}\n{proc.stdout}"
        )
    doc = json.loads(proc.stdout)
    assert "summary" in doc and "line_items" in doc and "packages" in doc
    assert doc["payment_lag_days"] == 30
    assert doc.get("calendar", {}).get("months")
    json_path = _SAMPLE / ".skill" / "cashflow_forecast.json"
    xlsx_path = _SAMPLE / "Skill4 Demo Project_现金流预测.xlsx"
    if not json_path.is_file() or not xlsx_path.is_file():
        raise SystemExit(f"Outputs missing under {_SAMPLE}")
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    ws = wb["现金流"]
    headers = [ws.cell(3, c).value for c in range(1, 9)]
    assert headers[:3] == ["序号", "采购包", "供应商"]
    assert "应请" in [ws.cell(3, c).value for c in range(1, 40)]
    assert "实请" in [ws.cell(3, c).value for c in range(1, 40)]
    assert ws.cell(4, 1).value == "01"
    assert ws.cell(5, 1).value == "总计"
    assert ws.cell(1, 8).value == "检查台账更新"
    assert "台账来源" in wb.sheetnames
    src_ws = wb["台账来源"]
    assert src_ws.cell(4, 1).value == "一键检查是否有更新"
    assert doc.get("sources")
    cmd = _SAMPLE / "检查现金流台账.command"
    if not cmd.is_file():
        raise SystemExit("check command missing")

    from check_sources import check

    first = check(_SAMPLE)
    assert first["fresh"] is True
    tracker = next(p for p in _SAMPLE.glob("*_procurement_packages_status_tracker.csv"))
    original = tracker.read_text(encoding="utf-8-sig")
    tracker.write_text(original + "\n", encoding="utf-8-sig")
    try:
        second = check(_SAMPLE)
        assert second["fresh"] is False
        assert any("Tracker" in str(c.get("label")) for c in second["changes"])
    finally:
        tracker.write_text(original, encoding="utf-8-sig")

    print(
        f"OK {_SAMPLE.name}: packages={len(doc['packages'])} "
        f"months={len(doc['calendar']['months'])} xlsx={xlsx_path.name}"
    )
    print("smoke_skill4: all passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
