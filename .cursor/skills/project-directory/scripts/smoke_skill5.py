#!/usr/bin/env python3
"""Smoke test Skill 5: group post, parse chat paste, write Excel."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

_SKILL_ROOT = Path(__file__).resolve().parents[1]
_SCRIPTS = _SKILL_ROOT / "scripts"
_SAMPLE = _SKILL_ROOT / "samples" / "demo_project"

sys.path.insert(0, str(_SCRIPTS))
from directory_store import parse_contact_blocks, split_pdir_sections  # noqa: E402


REPLY = """王工 16:02
[PDIR]
采购包：BMS
供应商：Schneider Electric (China) Co., Ltd.

---
姓名 / Name：王工
Title：项目经理
手机 / Mobile：13800138000
邮箱 / Email：wang@schneider.example.com
---
姓名 / Name：赵工
Title：现场负责人
手机 / Mobile：13900139000
邮箱 / Email：zhao@schneider.example.com
---
"""

QUOTED = """姓名 / Name：王工
Title：项目经理
手机 / Mobile：13800138000
邮箱 / Email：wang@schneider.example.com

-----Original Message-----
[PDIR]
采购包：BMS
供应商：Schneider

---
姓名 / Name：
Title：
手机 / Mobile：
邮箱 / Email：
---
"""


def run(script: str, extra: list[str]) -> dict:
    proc = subprocess.run(
        [sys.executable, str(_SCRIPTS / script), *extra],
        capture_output=True,
        text=True,
        check=False,
    )
    if proc.returncode != 0:
        raise SystemExit(f"{script} failed:\n{proc.stderr}\n{proc.stdout}")
    return json.loads(proc.stdout)


def reset_sample() -> None:
    skill = _SAMPLE / ".skill"
    for name in ("directory.json", "pending.json"):
        path = skill / name
        if path.exists():
            path.unlink()
    ddir = skill / "directory"
    if ddir.exists():
        shutil.rmtree(ddir)
    share = _SAMPLE / "项目档案" / "05_项目通讯录"
    if share.exists():
        shutil.rmtree(share)
    for leftover in _SAMPLE.glob("*_项目通讯录.xlsx"):
        leftover.unlink()


def main() -> int:
    people = parse_contact_blocks(REPLY)
    assert len(people) == 2, people
    assert people[0]["name"] == "王工"
    assert people[1]["email"] == "zhao@schneider.example.com"

    quoted = parse_contact_blocks(QUOTED)
    assert len(quoted) == 1, quoted
    assert quoted[0]["name"] == "王工"

    empty = parse_contact_blocks(
        "[PDIR]\n采购包：BMS\n---\n姓名 / Name：\nTitle：\n手机 / Mobile：\n邮箱 / Email：\n---\n"
    )
    assert empty == [], empty
    assert len(split_pdir_sections(REPLY)) == 1

    reset_sample()
    composed = run(
        "compose_group_post.py",
        ["--project-root", str(_SAMPLE), "--mark-posted"],
    )
    assert composed["channel"] == "wecom_group", composed
    assert Path(composed["group_message_path"]).is_file()
    assert Path(composed["fill_form_path"]).is_file()
    assert "[PDIR]" in composed["group_message"]
    assert "不必发邮件" in composed["group_message"]
    assert "smtp" not in composed["group_message"].lower()

    ingested = run(
        "ingest_directory_reply.py",
        ["--project-root", str(_SAMPLE), "--text", REPLY],
    )
    assert ingested["ok"] is True, ingested
    assert ingested["contacts_written"] == 2, ingested
    xlsx = Path(ingested["xlsx_path"])
    assert xlsx.is_file(), xlsx
    assert xlsx.name == "Skill5 Demo Project_项目通讯录.xlsx", xlsx.name

    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    ws = wb["项目通讯录"]
    headers = [ws.cell(2, c).value for c in range(1, 9)]
    assert headers == ["序号", "采购包", "供应商", "姓名", "Title", "手机", "邮箱", "更新时间"], headers
    assert ws.cell(3, 4).value == "王工"
    ws2 = wb["收集进度"]
    statuses = {ws2.cell(r, 2).value: ws2.cell(r, 4).value for r in range(3, 5)}
    assert statuses["BMS"] == "已收录", statuses
    assert statuses["HVAC"] == "已发群待回", statuses

    print(f"OK {_SAMPLE.name}: contacts=2 xlsx={xlsx.name} channel=wecom_group")
    print("smoke_skill5: all passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
