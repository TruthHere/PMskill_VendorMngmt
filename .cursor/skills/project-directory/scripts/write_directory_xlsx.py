#!/usr/bin/env python3
"""Write Project Directory workbook: contacts + collection progress."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

_SCRIPTS = Path(__file__).resolve().parent
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))

from directory_store import (  # noqa: E402
    load_store,
    project_meta,
    request_for,
    save_store,
    status_label,
    xlsx_path_for,
)

CONTACT_HEADERS = ["序号", "采购包", "供应商", "姓名", "Title", "手机", "邮箱", "更新时间"]
PROGRESS_HEADERS = [
    "序号",
    "采购包",
    "供应商",
    "收集状态",
    "群消息日期",
    "回执日期",
    "已收录人数",
    "备注",
]


def _theme() -> dict[str, str]:
    return {
        "ink": "0F172A",
        "ink_muted": "64748B",
        "header_bg": "334155",
        "header_text": "FFFFFF",
        "zebra": "F8FAFC",
        "border": "E2E8F0",
        "ok_bg": "DBEAFE",
        "ok_text": "1D4ED8",
        "warn_bg": "FEF3C7",
        "warn_text": "92400E",
        "bad_bg": "FEE2E2",
        "bad_text": "991B1B",
        "font": "Microsoft YaHei",
    }


def _day(iso: str) -> str:
    text = (iso or "").strip()
    if not text:
        return ""
    return text[:10]


def write_directory_xlsx(project_root: Path, store: dict[str, Any] | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet

    meta = project_meta(project_root)
    store = store if store is not None else load_store(project_root)
    project_name = store.get("project_name") or meta["project_name"]
    path = xlsx_path_for(meta["root"], str(project_name))
    theme = _theme()
    thin = Side(style="thin", color=theme["border"])
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731
    fonts = {
        "title": Font(name=theme["font"], size=12, bold=True, color=theme["ink"]),
        "header": Font(name=theme["font"], size=10, bold=True, color=theme["header_text"]),
        "body": Font(name=theme["font"], size=10, color=theme["ink"]),
        "ok": Font(name=theme["font"], size=10, color=theme["ok_text"]),
        "warn": Font(name=theme["font"], size=10, color=theme["warn_text"]),
        "bad": Font(name=theme["font"], size=10, color=theme["bad_text"]),
    }
    fills = {
        "header": fill(theme["header_bg"]),
        "zebra": fill(theme["zebra"]),
        "ok": fill(theme["ok_bg"]),
        "warn": fill(theme["warn_bg"]),
        "bad": fill(theme["bad_bg"]),
    }
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _paint_header(ws: Worksheet, headers: list[str], row: int = 2) -> None:
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row, col, title)
            cell.font = fonts["header"]
            cell.fill = fills["header"]
            cell.alignment = center
            cell.border = border
        ws.row_dimensions[row].height = 22
        ws.freeze_panes = f"A{row + 1}"

    def _body(cell, *, align=left, font=None, fill_key: str | None = None) -> None:
        cell.font = font or fonts["body"]
        cell.alignment = align
        cell.border = border
        if fill_key:
            cell.fill = fills[fill_key]

    wb = Workbook()
    ws = wb.active
    ws.title = "项目通讯录"
    updated = _day(str(store.get("updated_at") or ""))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(CONTACT_HEADERS))
    title = ws.cell(1, 1, f"项目通讯录 · {project_name} · 更新于 {updated or '—'}")
    title.font = fonts["title"]
    ws.row_dimensions[1].height = 24
    _paint_header(ws, CONTACT_HEADERS)

    contacts = list(store.get("contacts") or [])
    contacts.sort(key=lambda c: (str(c.get("seq") or ""), str(c.get("package_name") or ""), str(c.get("name") or "")))
    for i, person in enumerate(contacts):
        row = i + 3
        values = [
            person.get("seq") or "",
            person.get("package_name") or "",
            person.get("supplier_name") or "",
            person.get("name") or "",
            person.get("title") or "",
            person.get("phone") or "",
            person.get("email") or "",
            _day(str(person.get("updated_at") or "")),
        ]
        zebra = i % 2 == 1
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            align = center if col in {1, 8} else left
            _body(cell, align=align, fill_key="zebra" if zebra else None)
            if col == 6:
                cell.number_format = "@"
        ws.row_dimensions[row].height = 22

    widths = {"A": 8, "B": 16, "C": 28, "D": 12, "E": 16, "F": 16, "G": 32, "H": 14}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    last_contact = max(2, ws.max_row)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(CONTACT_HEADERS))}{last_contact}"

    ws2 = wb.create_sheet("收集进度")
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(PROGRESS_HEADERS))
    ws2.cell(1, 1, f"收集进度 · {project_name}").font = fonts["title"]
    ws2.row_dimensions[1].height = 24
    _paint_header(ws2, PROGRESS_HEADERS)

    packages = meta["packages"]
    counts: dict[str, int] = {}
    for person in contacts:
        key = str(person.get("package_name") or "")
        counts[key] = counts.get(key, 0) + 1

    for i, pkg in enumerate(packages):
        row = i + 3
        name = str(pkg.get("package_name") or "")
        req = request_for(store, name) or {}
        status = str(req.get("status") or "ready")
        label = status_label(status)
        note = str(req.get("error") or "")
        if status == "ready" and not note:
            note = "把群消息发到项目群；填写页丢进群文件"
        posted = str(req.get("posted_at") or req.get("sent_at") or req.get("ready_at") or "")
        values = [
            pkg.get("seq") or "",
            name,
            pkg.get("supplier_name") or pkg.get("supplier_short") or "",
            label,
            _day(posted),
            _day(str(req.get("replied_at") or "")),
            counts.get(name, 0),
            note,
        ]
        style = "ok" if status == "replied" else "warn" if status in {"posted", "sent", "ready"} else None
        font = fonts["ok"] if style == "ok" else fonts["warn"] if style == "warn" else fonts["body"]
        zebra = i % 2 == 1
        for col, value in enumerate(values, start=1):
            cell = ws2.cell(row, col, value)
            align = center if col in {1, 4, 5, 6, 7} else left
            fill_key = style if col == 4 and style else ("zebra" if zebra else None)
            _body(cell, align=align, font=font if col == 4 else fonts["body"], fill_key=fill_key)
        ws2.row_dimensions[row].height = 22

    widths2 = {"A": 8, "B": 16, "C": 28, "D": 16, "E": 14, "F": 14, "G": 12, "H": 42}
    for letter, width in widths2.items():
        ws2.column_dimensions[letter].width = width
    last_prog = max(2, ws2.max_row)
    ws2.auto_filter.ref = f"A2:{get_column_letter(len(PROGRESS_HEADERS))}{last_prog}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    store["xlsx_path"] = str(path)
    store["project_name"] = project_name
    save_store(project_root, store)
    return path


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--project-root", required=True)
    args = ap.parse_args()
    root = Path(args.project_root).expanduser().resolve()
    path = write_directory_xlsx(root)
    print(json.dumps({"ok": True, "xlsx_path": str(path)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
