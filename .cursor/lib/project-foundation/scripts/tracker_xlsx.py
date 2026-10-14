#!/usr/bin/env python3
"""Tracker xlsx: create once with theme, then incremental data sync."""

from __future__ import annotations

import math
import re
import warnings
from pathlib import Path
from typing import Any, Callable

from common import (
    MILESTONE_TO_CSV,
    PAYMENT_CSV_HEADERS,
    TRACKER_FIELDNAMES,
    _milestone_column_range,
    _payment_column_range,
    format_milestone_date,
    pkg_to_csv_row,
)

TRACKER_XLSX_SCHEMA_VERSION = 2
GROUP_ROW = 2
HEADER_ROW = 3
DATA_START = HEADER_ROW + 1


def _theme() -> dict[str, Any]:
    return {
        "font": "Microsoft YaHei",
        "ink": "0F172A",
        "ink_muted": "64748B",
        "header_bg": "334155",
        "header_text": "FFFFFF",
        "group_bg": "F1F5F9",
        "group_text": "475569",
        "zebra": "F8FAFC",
        "pay_bg": "FFFBEB",
        "pay_zebra": "FEF3C7",
        "pay_hl": "FFF7ED",
        "pay_group_text": "92400E",
        "milestone_hdr": "475569",
        "milestone_done_bg": "DBEAFE",
        "milestone_done_text": "1D4ED8",
        "border": "E2E8F0",
        "col_widths": {
            "序号": 5,
            "采购包名称": 16,
            "供应商名称": 20,
            "供应商简称": 11,
            "对接人邮箱": 26,
            "采购包类型": 9,
            "金额不含税": 13,
            "金额含税": 13,
            "金额备注": 18,
            "交期摘要": 22,
            "已请金额": 13,
            "待请款金额": 13,
            "下一付款节点": 18,
            "请款进度": 22,
            "RFP 发出": 11,
            "技术方案已确认": 13,
            "商务报价已确认": 13,
            "PO 已发起": 11,
            "已结算": 10,
        },
        "wrap_headers": {
            "采购包名称",
            "供应商名称",
            "供应商简称",
            "对接人邮箱",
            "金额备注",
            "交期摘要",
            "下一付款节点",
            "请款进度",
        },
        "center_headers": {"序号", "采购包类型", *MILESTONE_TO_CSV.values()},
        "right_headers": {"金额不含税", "金额含税", "已请金额", "待请款金额"},
    }


def _styles(theme: dict[str, Any]):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color=theme["border"])
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731

    def cell_alignment(header: str) -> Alignment:
        wrap = header in theme["wrap_headers"]
        if header in theme["center_headers"]:
            return Alignment(horizontal="center", vertical="center", wrap_text=wrap)
        if header in theme["right_headers"]:
            return Alignment(horizontal="right", vertical="center", wrap_text=False)
        return Alignment(
            horizontal="left",
            vertical="top" if wrap else "center",
            wrap_text=wrap,
        )

    fonts = {
        "group": Font(name=theme["font"], size=10, bold=True, color=theme["group_text"]),
        "header": Font(name=theme["font"], size=10, bold=True, color=theme["header_text"]),
        "body": Font(name=theme["font"], size=10, color=theme["ink"]),
        "meta": Font(name=theme["font"], size=10, color=theme["ink_muted"]),
        "pay_group": Font(
            name=theme["font"], size=10, bold=True, color=theme["pay_group_text"]
        ),
        "pay_hl": Font(name=theme["font"], size=10, bold=True, color=theme["ink"]),
        "milestone_done": Font(
            name=theme["font"], size=10, bold=True, color=theme["milestone_done_text"]
        ),
    }
    fills = {
        "header": fill(theme["header_bg"]),
        "group": fill(theme["group_bg"]),
        "zebra": fill(theme["zebra"]),
        "pay": fill(theme["pay_bg"]),
        "pay_zebra": fill(theme["pay_zebra"]),
        "pay_hl": fill(theme["pay_hl"]),
        "milestone_hdr": fill(theme["milestone_hdr"]),
        "milestone_done": fill(theme["milestone_done_bg"]),
    }
    return border_all, cell_alignment, fonts, fills


def _display_width(text: str) -> int:
    return sum(2 if ord(ch) > 127 else 1 for ch in str(text or ""))


def _row_height_for_values(values: dict[str, str], theme: dict[str, Any]) -> float:
    col_widths = theme["col_widths"]
    wrap_headers = theme["wrap_headers"]
    max_lines = 1
    for header in wrap_headers:
        text = values.get(header, "")
        if not text:
            continue
        capacity = max(6, int(col_widths.get(header, 12) * 1.45))
        lines = max(1, math.ceil(_display_width(text) / capacity))
        max_lines = max(max_lines, lines)
    return min(120.0, max(22.0, max_lines * 15.0 + 8.0))


def _write_cell_value(cell, header: str, value: str) -> None:
    if header in ("金额不含税", "金额含税", "已请金额", "待请款金额") and value:
        raw = str(value).strip()
        # Accept plain decimals or currency-prefixed strings like ¥12,345.67
        m = re.search(r"([\d,]+(?:\.\d+)?)", raw.replace(" ", ""))
        if m:
            try:
                cell.value = float(m.group(1).replace(",", ""))
                cell.number_format = "#,##0.00"
                return
            except ValueError:
                pass
    cell.value = value if value else None


def _apply_milestone_style(cell, header: str, value: str, fonts, fills) -> None:
    if header not in MILESTONE_TO_CSV.values():
        return
    if value:
        cell.fill = fills["milestone_done"]
        cell.font = fonts["milestone_done"]


def _apply_payment_style(cell, header: str, value: str, is_zebra: bool, fonts, fills) -> None:
    if header not in PAYMENT_CSV_HEADERS:
        return
    if header == "下一付款节点" and value:
        cell.fill = fills["pay_hl"]
        cell.font = fonts["pay_hl"]
        return
    cell.fill = fills["pay_zebra"] if is_zebra else fills["pay"]


def _read_header_row(ws) -> list[str]:
    return [str(ws.cell(HEADER_ROW, c).value or "").strip() for c in range(1, len(TRACKER_FIELDNAMES) + 1)]


def _schema_matches(ws) -> bool:
    return _read_header_row(ws) == TRACKER_FIELDNAMES


def _package_name_col() -> int:
    return TRACKER_FIELDNAMES.index("采购包名称") + 1


def _map_existing_rows(ws) -> dict[str, int]:
    col = _package_name_col()
    mapping: dict[str, int] = {}
    for row in range(DATA_START, ws.max_row + 1):
        name = str(ws.cell(row, col).value or "").strip()
        if name:
            mapping[name] = row
    return mapping


def _update_meta_row(ws, updated_at: str, fonts, project_name: str = "") -> None:
    updated_date = format_milestone_date(updated_at) or updated_at[:10]
    prefix = f"{project_name} · " if project_name else ""
    cell = ws.cell(1, 1)
    cell.value = f"{prefix}最后更新：{updated_date}"
    if not cell.font or cell.font.name is None:
        cell.font = fonts["meta"]


def _apply_data_row_style(
    ws,
    row_idx: int,
    row: dict[str, str],
    *,
    is_zebra: bool,
    theme: dict[str, Any],
    border_all,
    cell_alignment: Callable[[str], Any],
    fonts,
    fills,
) -> None:
    milestone_cols = {TRACKER_FIELDNAMES.index(c) + 1 for c in MILESTONE_TO_CSV.values()}
    payment_cols = {TRACKER_FIELDNAMES.index(c) + 1 for c in PAYMENT_CSV_HEADERS}
    for col_idx, header in enumerate(TRACKER_FIELDNAMES, start=1):
        value = row.get(header, "")
        cell = ws.cell(row_idx, col_idx)
        cell.border = border_all
        cell.alignment = cell_alignment(header)
        if header in MILESTONE_TO_CSV.values() and value:
            cell.font = fonts["milestone_done"]
        elif header == "下一付款节点" and value:
            cell.font = fonts["pay_hl"]
        else:
            cell.font = fonts["body"]
        if col_idx in payment_cols:
            _apply_payment_style(cell, header, value, is_zebra, fonts, fills)
        elif is_zebra and col_idx not in milestone_cols:
            cell.fill = fills["zebra"]
        _write_cell_value(cell, header, value)
        _apply_milestone_style(cell, header, value, fonts, fills)
    ws.row_dimensions[row_idx].height = _row_height_for_values(row, theme)


def _sync_data_rows(ws, packages: list[dict[str, Any]], theme: dict[str, Any], fonts, fills) -> None:
    """只更新数据格的值；已有行的样式/行高/列宽尽量保留。"""
    border_all, cell_alignment, _, _ = _styles(theme)
    existing = _map_existing_rows(ws)
    desired_names = [str(p.get("package_name") or "").strip() for p in packages]
    desired_set = {n for n in desired_names if n}

    for pkg in packages:
        row_data = pkg_to_csv_row(pkg)
        name = str(pkg.get("package_name") or "").strip()
        if not name:
            continue
        if name in existing:
            row_idx = existing[name]
            is_zebra = (row_idx - DATA_START) % 2 == 1
            for col_idx, header in enumerate(TRACKER_FIELDNAMES, start=1):
                cell = ws.cell(row_idx, col_idx)
                value = row_data.get(header, "")
                _write_cell_value(cell, header, value)
                _apply_payment_style(cell, header, value, is_zebra, fonts, fills)
                _apply_milestone_style(cell, header, value, fonts, fills)
                if cell.alignment is None or cell.alignment.horizontal is None:
                    cell.alignment = cell_alignment(header)
                if cell.border is None or (
                    cell.border.left is None and cell.border.right is None
                ):
                    cell.border = border_all
        else:
            row_idx = ws.max_row + 1 if ws.max_row >= DATA_START else DATA_START
            is_zebra = (row_idx - DATA_START) % 2 == 1
            _apply_data_row_style(
                ws,
                row_idx,
                row_data,
                is_zebra=is_zebra,
                theme=theme,
                border_all=border_all,
                cell_alignment=cell_alignment,
                fonts=fonts,
                fills=fills,
            )
            existing[name] = row_idx

    for name, row_idx in sorted(existing.items(), key=lambda x: x[1], reverse=True):
        if name not in desired_set:
            ws.delete_rows(row_idx, 1)

    if ws.max_row >= DATA_START:
        from openpyxl.utils import get_column_letter

        last_col = len(TRACKER_FIELDNAMES)
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(last_col)}{ws.max_row}"


def _create_workbook(
    packages: list[dict[str, Any]], updated_at: str, project_name: str = ""
) -> Any:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    theme = _theme()
    border_all, cell_alignment, fonts, fills = _styles(theme)
    wb = Workbook()
    ws = wb.active
    ws.title = "采购包 Tracker"
    last_col = len(TRACKER_FIELDNAMES)

    _update_meta_row(ws, updated_at, fonts, project_name)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    pay_first, pay_last = _payment_column_range()
    ms_first, ms_last = _milestone_column_range()
    group_ranges = (("请款", pay_first, pay_last), ("状态", ms_first, ms_last))
    for label, first, last in group_ranges:
        if last < first:
            continue
        ws.merge_cells(
            start_row=GROUP_ROW,
            start_column=first + 1,
            end_row=GROUP_ROW,
            end_column=last + 1,
        )
        group_cell = ws.cell(GROUP_ROW, first + 1, label)
        group_cell.alignment = Alignment(horizontal="center", vertical="center")
        if label == "请款":
            group_cell.font = fonts["pay_group"]
            group_cell.fill = fills["pay"]
        else:
            group_cell.font = fonts["group"]
            group_cell.fill = fills["group"]
        group_cell.border = border_all

    pay_cols = set(range(pay_first + 1, pay_last + 2))
    ms_cols = set(range(ms_first + 1, ms_last + 2))
    for col_idx in range(1, last_col + 1):
        cell = ws.cell(GROUP_ROW, col_idx)
        cell.fill = fills["pay"] if col_idx in pay_cols else fills["group"]
        cell.border = border_all
        if col_idx not in pay_cols and col_idx not in ms_cols:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(GROUP_ROW, pay_first + 1).font = fonts["pay_group"]
    ws.cell(GROUP_ROW, pay_first + 1).fill = fills["pay"]

    for col_idx, header in enumerate(TRACKER_FIELDNAMES, start=1):
        cell = ws.cell(HEADER_ROW, col_idx, header)
        cell.font = fonts["header"]
        if header in MILESTONE_TO_CSV.values() or header in PAYMENT_CSV_HEADERS:
            cell.fill = fills["milestone_hdr"]
        else:
            cell.fill = fills["header"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_all

    ws.row_dimensions[GROUP_ROW].height = 24
    ws.row_dimensions[HEADER_ROW].height = 32

    for row_idx, pkg in enumerate(packages, start=DATA_START):
        row = pkg_to_csv_row(pkg)
        is_zebra = (row_idx - DATA_START) % 2 == 1
        _apply_data_row_style(
            ws,
            row_idx,
            row,
            is_zebra=is_zebra,
            theme=theme,
            border_all=border_all,
            cell_alignment=cell_alignment,
            fonts=fonts,
            fills=fills,
        )

    for col_idx, header in enumerate(TRACKER_FIELDNAMES, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = theme["col_widths"].get(header, 12)

    ws.freeze_panes = ws.cell(DATA_START, 1)
    if packages:
        ws.auto_filter.ref = (
            f"A{HEADER_ROW}:{get_column_letter(last_col)}{DATA_START + len(packages) - 1}"
        )
    ws.sheet_properties.tabColor = theme["header_bg"]
    _assert_layout(ws, cell_alignment, len(packages))
    return wb


def _assert_layout(ws, cell_alignment: Callable[[str], Any], package_count: int) -> None:
    if package_count == 0 or ws.max_row < DATA_START:
        return
    for col_idx, header in enumerate(TRACKER_FIELDNAMES, start=1):
        cell = ws.cell(DATA_START, col_idx)
        expected = cell_alignment(header)
        if cell.alignment.horizontal != expected.horizontal:
            warnings.warn(
                f"Tracker xlsx 对齐不一致：列「{header}」"
                f" 期望 {expected.horizontal}，实际 {cell.alignment.horizontal}",
                stacklevel=2,
            )


def save_tracker_xlsx(
    path: Path,
    packages: list[dict[str, Any]],
    updated_at: str,
    project_name: str = "",
) -> None:
    """
    xlsx 写入策略：
    - 文件不存在 / 表头与当前 schema 不一致 → 全量创建（含主题样式）
    - 否则 → 增量同步：只改数据格 + 最后更新 + 里程碑完成色；保留列宽/行高/手动排版
    """
    path = path.expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    theme = _theme()
    _, _, fonts, fills = _styles(theme)

    if path.is_file():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path)
            ws = wb.active
            if _schema_matches(ws):
                _update_meta_row(ws, updated_at, fonts, project_name)
                _sync_data_rows(ws, packages, theme, fonts, fills)
                wb.save(path)
                return
        except OSError:
            pass
        except Exception as exc:
            warnings.warn(f"Tracker xlsx 增量更新失败，将全量重建：{exc}", stacklevel=2)

    wb = _create_workbook(packages, updated_at, project_name)
    wb.save(path)
