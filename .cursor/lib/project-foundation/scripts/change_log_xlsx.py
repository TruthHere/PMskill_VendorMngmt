#!/usr/bin/env python3
"""Write human-facing Change Log Excel at project root (CSV stays in .skill/)."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def save_change_log_xlsx(
    xlsx_path: Path,
    rows: list[dict[str, str]],
    fieldnames: list[str],
    *,
    project_name: str = "",
    updated_at: str = "",
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Change Log"

    ink = "0F172A"
    hdr_bg = "334155"
    hdr_fg = "FFFFFF"
    zebra = "F8FAFC"
    thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    meta = f"Change Log · {project_name or 'Project'} · updated {updated_at or '—'}"
    ws.cell(1, 1, meta)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(fieldnames)))
    ws.cell(1, 1).font = Font(name="Microsoft YaHei", size=11, color="64748B")

    header_fill = PatternFill("solid", fgColor=hdr_bg)
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color=hdr_fg)
    body_font = Font(name="Microsoft YaHei", size=10, color=ink)
    for col, name in enumerate(fieldnames, start=1):
        cell = ws.cell(2, col, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    zebra_fill = PatternFill("solid", fgColor=zebra)
    for r_idx, row in enumerate(rows):
        excel_r = r_idx + 3
        for c_idx, name in enumerate(fieldnames, start=1):
            cell = ws.cell(excel_r, c_idx, row.get(name, ""))
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=name in ("变更内容", "详细说明", "备注"))
            cell.border = thin
            if r_idx % 2 == 1:
                cell.fill = zebra_fill

    widths = {
        "变更序号": 12,
        "采购包名称": 16,
        "供应商": 14,
        "变更内容": 28,
        "详细说明": 36,
        "预估金额": 12,
        "最终确认金额": 14,
        "时间影响": 14,
        "当前状态": 14,
        "归档路径": 28,
        "备注": 22,
    }
    for c_idx, name in enumerate(fieldnames, start=1):
        ws.column_dimensions[ws.cell(2, c_idx).column_letter].width = widths.get(name, 12)

    ws.freeze_panes = "A3"
    ws.row_dimensions[2].height = 28
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    return xlsx_path


def sync_change_log_xlsx_from_csv(project_root: Path, csv_path: Path) -> Path | None:
    """读 CSV 后写出根目录 Change Log.xlsx。"""
    import csv
    import io
    import sys

    scripts = Path(__file__).resolve().parent
    if str(scripts) not in sys.path:
        sys.path.insert(0, str(scripts))
    from common import (  # noqa: WPS433
        load_json,
        now_iso,
        resolve_change_log_xlsx_path,
        system_file,
        PROJECT_JSON,
    )

    root = project_root.expanduser().resolve()
    xlsx = resolve_change_log_xlsx_path(root)
    if xlsx is None:
        return None
    raw = csv_path.read_text(encoding="utf-8-sig")
    lines = [ln for ln in raw.splitlines() if ln.strip() and not ln.lstrip().startswith("#")]
    if not lines:
        fieldnames: list[str] = []
        rows: list[dict[str, str]] = []
    else:
        reader = csv.DictReader(io.StringIO("\n".join(lines)))
        fieldnames = list(reader.fieldnames or [])
        rows = [{k: (v or "") for k, v in row.items() if k} for row in reader]
    project = load_json(system_file(root, PROJECT_JSON), {}) or {}
    return save_change_log_xlsx(
        xlsx,
        rows,
        fieldnames,
        project_name=str(project.get("project_name") or ""),
        updated_at=now_iso(),
    )
