#!/usr/bin/env python3
"""Write Project-like Excel: left task table + right Gantt bars."""

from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

from readers import TaskRow, iter_timeline, timeline_bounds

INK = "0F172A"
HEADER_BG = "334155"
HEADER_TEXT = "FFFFFF"
MONTH_BG = "475569"
SUMMARY_ROW = "EEF2F6"
ZEBRA = "F8FAFC"
BAR_TASK = "5B8DEF"
BAR_SUMMARY = "334155"
BAR_MILESTONE = "0F172A"
BAR_CRITICAL = "DC2626"
GRID = "E2E8F0"
META_BG = "F1F5F9"

# Human Project view columns (package stays in structured export only)
LEFT_HEADERS = [
    "ID",
    "WBS",
    "任务名称",
    "工期",
    "开始",
    "结束",
    "前置任务",
    "%完成",
]


def _fill(hex_rgb: str):
    from openpyxl.styles import PatternFill

    return PatternFill(start_color=hex_rgb, end_color=hex_rgb, fill_type="solid")


def _font(bold: bool = False, color: str = INK, size: int = 11):
    from openpyxl.styles import Font

    return Font(name="Calibri", bold=bold, color=color, size=size)


def _align(h: str = "left", v: str = "center", indent: int = 0):
    from openpyxl.styles import Alignment

    return Alignment(horizontal=h, vertical=v, indent=indent, wrap_text=False)


def _border_light():
    from openpyxl.styles import Border, Side

    side = Side(style="thin", color=GRID)
    return Border(left=side, right=side, top=side, bottom=side)


def _no_border():
    from openpyxl.styles import Border

    return Border()


def write_project_view(
    tasks: list[TaskRow],
    out_path: Path,
    *,
    source_name: str = "",
    project_name: str = "",
    gantt_unit: str = "week",
) -> Path:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    if gantt_unit not in {"day", "week"}:
        raise ValueError("gantt_unit 须为 day 或 week")

    wb = Workbook()
    ws = wb.active
    ws.title = "Project View"

    lo, hi = timeline_bounds(tasks, gantt_unit)
    ticks = iter_timeline(lo, hi, gantt_unit)
    left_n = len(LEFT_HEADERS)
    gantt_start_col = left_n + 1

    # Row 1: metadata across left + gantt
    meta_end = max(left_n, gantt_start_col + max(len(ticks) - 1, 0))
    meta = (
        f"项目: {project_name or '—'}  |  源: {source_name or '—'}  |  "
        f"导出: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  甘特: {gantt_unit}"
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=meta_end)
    c = ws.cell(1, 1, meta)
    c.fill = _fill(META_BG)
    c.font = _font(False, "64748B", 10)
    c.alignment = _align("left", "center")
    ws.row_dimensions[1].height = 18

    # Row 2: month band (gantt only) + blank left header area styled
    for col in range(1, left_n + 1):
        cell = ws.cell(2, col, "")
        cell.fill = _fill(HEADER_BG)
        cell.border = _border_light()

    _write_month_band(ws, ticks, gantt_start_col, gantt_unit)

    # Row 3: column headers + week/day labels
    for col, title in enumerate(LEFT_HEADERS, start=1):
        cell = ws.cell(3, col, title)
        cell.fill = _fill(HEADER_BG)
        cell.font = _font(True, HEADER_TEXT, 10)
        cell.alignment = _align("center", "center")
        cell.border = _border_light()

    for i, tick in enumerate(ticks):
        col = gantt_start_col + i
        label = f"{tick.month}/{tick.day}"
        cell = ws.cell(3, col, label)
        cell.fill = _fill(HEADER_BG)
        cell.font = _font(True, HEADER_TEXT, 8)
        cell.alignment = _align("center", "center")
        cell.border = _border_light()
        ws.column_dimensions[get_column_letter(col)].width = 3.6 if gantt_unit == "day" else 4.5

    data_start_row = 4

    for r_i, task in enumerate(tasks):
        row = data_start_row + r_i
        zebra = r_i % 2 == 1
        base_fill = SUMMARY_ROW if task.is_summary else (ZEBRA if zebra else "FFFFFF")

        values = [
            task.task_id,
            task.wbs,
            task.name,
            task.duration,
            task.start.isoformat() if task.start else "",
            task.finish.isoformat() if task.finish else "",
            task.predecessors,
            int(task.percent_complete) if task.percent_complete else 0,
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row, col, val)
            cell.fill = _fill(base_fill)
            cell.border = _border_light()
            cell.font = _font(bold=task.is_summary, color=INK, size=10)
            if col in {1, 4, 8}:
                cell.alignment = _align("center", "center")
            elif col in {5, 6}:
                cell.alignment = _align("center", "center")
            elif col == 2:
                cell.alignment = _align("left", "center")
            elif col == 3:
                cell.alignment = _align("left", "center", indent=max(0, task.outline_level - 1))
            else:
                cell.alignment = _align("left", "center")

        for i, tick in enumerate(ticks):
            col = gantt_start_col + i
            cell = ws.cell(row, col, "")
            cell.border = _no_border()
            hit = _bar_hit(task, tick, gantt_unit)
            if hit == "milestone":
                cell.value = "◆"
                cell.fill = _fill(BAR_CRITICAL if task.is_critical else BAR_MILESTONE)
                cell.font = _font(True, "FFFFFF", 8)
                cell.alignment = _align("center", "center")
            elif hit == "bar":
                if task.is_critical:
                    color = BAR_CRITICAL
                elif task.is_summary:
                    color = BAR_SUMMARY
                else:
                    color = BAR_TASK
                cell.fill = _fill(color)
            # else: leave default (no fill / no border) — clean gantt

        ws.row_dimensions[row].height = 17

    widths = [5, 10, 38, 7, 11, 11, 14, 7]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze: meta+headers above, all left columns locked → gantt scrolls
    freeze_col = get_column_letter(gantt_start_col)
    ws.freeze_panes = f"{freeze_col}{data_start_row}"
    ws.print_title_rows = "1:3"

    _write_legend_sheet(wb)

    # Point to legend from meta row
    c.value = f"{meta}  |  色例见图例工作表"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _write_legend_sheet(wb) -> None:
    """Explain gantt bar colors for PM readers."""
    from openpyxl.utils import get_column_letter

    # Insert as second sheet
    if "图例" in wb.sheetnames:
        ws = wb["图例"]
        wb.remove(ws)
    ws = wb.create_sheet("图例", 1)

    ws["A1"] = "甘特色块说明"
    ws["A1"].font = _font(True, INK, 14)
    ws.merge_cells("A1:C1")

    ws["A2"] = "打开「Project View」时，右侧色块含义如下（与 MS Project 甘特语义接近）。"
    ws["A2"].font = _font(False, "64748B", 10)
    ws.merge_cells("A2:C2")

    headers = ["色块", "含义", "说明"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(4, col, h)
        cell.fill = _fill(HEADER_BG)
        cell.font = _font(True, HEADER_TEXT, 10)
        cell.alignment = _align("center", "center")
        cell.border = _border_light()

    rows = [
        (BAR_TASK, "", "普通任务", "叶子任务的计划工期条（非关键路径）"),
        (BAR_SUMMARY, "", "摘要任务", "汇总/父级任务条（WBS 上级），对应 Project 的 Summary"),
        (BAR_MILESTONE, "◆", "里程碑", "零工期节点；黑底白◆，对应 Project 的 Milestone"),
        (BAR_CRITICAL, "", "关键路径任务", "落在关键路径上的任务条（Project Critical）"),
        (BAR_CRITICAL, "◆", "关键路径里程碑", "既是里程碑又在关键路径上，红底白◆"),
    ]

    for i, (color, mark, meaning, note) in enumerate(rows):
        row = 5 + i
        swatch = ws.cell(row, 1, mark or "")
        swatch.fill = _fill(color)
        swatch.font = _font(True, "FFFFFF", 11)
        swatch.alignment = _align("center", "center")
        swatch.border = _border_light()

        m = ws.cell(row, 2, meaning)
        m.font = _font(True, INK, 10)
        m.alignment = _align("left", "center")
        m.border = _border_light()

        n = ws.cell(row, 3, note)
        n.font = _font(False, INK, 10)
        n.alignment = _align("left", "center")
        n.border = _border_light()
        ws.row_dimensions[row].height = 22

    ws["A11"] = "补充"
    ws["A11"].font = _font(True, INK, 11)
    ws["A12"] = "· 左侧任务名称加粗 + 浅灰行底 = 摘要任务（与甘特深色条对应）。"
    ws["A13"] = "· 空单元格无色、无边框，仅表示该周/日没有工期覆盖。"
    ws["A14"] = "· 关键路径以源计划 Critical 字段为准；本表不重新计算关键路径。"
    for r in range(12, 15):
        ws.cell(r, 1).font = _font(False, "64748B", 10)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 52


def _write_month_band(ws, ticks: list[date], gantt_start_col: int, unit: str) -> None:
    if not ticks:
        return
    # Group consecutive ticks by (year, month)
    groups: list[tuple[str, int, int]] = []  # label, start_idx, end_idx inclusive
    cur_key = (ticks[0].year, ticks[0].month)
    start_i = 0
    for i, tick in enumerate(ticks):
        key = (tick.year, tick.month)
        if key != cur_key:
            label = f"{cur_key[0]}-{cur_key[1]:02d}"
            groups.append((label, start_i, i - 1))
            cur_key = key
            start_i = i
    groups.append((f"{cur_key[0]}-{cur_key[1]:02d}", start_i, len(ticks) - 1))

    for label, a, b in groups:
        c0 = gantt_start_col + a
        c1 = gantt_start_col + b
        if c1 > c0:
            ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c1)
        cell = ws.cell(2, c0, label)
        cell.fill = _fill(MONTH_BG)
        cell.font = _font(True, HEADER_TEXT, 9)
        cell.alignment = _align("center", "center")
        cell.border = _border_light()
        for col in range(c0, c1 + 1):
            ws.cell(2, col).fill = _fill(MONTH_BG)
            ws.cell(2, col).border = _border_light()


def _bar_hit(task: TaskRow, tick: date, unit: str) -> Optional[str]:
    if not task.start and not task.finish:
        return None
    start = task.start or task.finish
    finish = task.finish or task.start
    assert start is not None and finish is not None

    if task.is_milestone or start == finish:
        if unit == "day":
            return "milestone" if tick == start else None
        return "milestone" if tick <= start <= tick + timedelta(days=6) else None

    if unit == "day":
        return "bar" if start <= tick <= finish else None

    window_end = tick + timedelta(days=6)
    if finish < tick or start > window_end:
        return None
    return "bar"
