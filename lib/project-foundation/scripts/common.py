#!/usr/bin/env python3
"""Shared utilities for project-foundation (Tracker / .skill / paths / amounts)."""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

CST = timezone(timedelta(hours=8))

PACKAGE_SUBDIRS_ALL = [
    "01_合同报价",
    "02_图纸",
    "03_请款记录",
    "04_变更洽商",
    "05_验收记录",
    "06_结算资料",
    "07_往来文件",
]
PACKAGE_SUBDIR_GC = "08_产值申报"
PACKAGE_SUBDIR_CONSULTANT = "09_交付物"

CARD_FILENAME = "履约条件卡片.html"
CARD_DATA_FILENAME = "card-data.json"
# 给人看的 Excel 在项目根；机读 CSV 在 .skill/（隐藏）
TRACKER_XLSX_SUFFIX = "_Status Tracker.xlsx"
TRACKER_CSV_NAME = "status_tracker.csv"  # 位于 .skill/
TRACKER_CSV_LEGACY = "tracker.csv"
TRACKER_CSV_LEGACY_SUFFIX = "_procurement_packages_status_tracker.csv"
CHANGE_LOG_XLSX_SUFFIX = "_Change Log.xlsx"
CHANGE_LOG_CSV_NAME = "change_log.csv"  # 位于 .skill/
CHANGE_LOG_CSV_LEGACY_SUFFIX = "_change_log.csv"
# 兼容旧常量名（部分脚本仍引用）
TRACKER_SUFFIX = TRACKER_CSV_LEGACY_SUFFIX
SYSTEM_DIR = ".skill"
LEGACY_SYSTEM_DIRS = ("_系统",)
PROJECT_ARCHIVE_DIR = "项目档案"
LEGACY_ARCHIVE_DIRS = ("_项目档案",)
BACKGROUND_MD = "项目背景.md"
PROJECT_JSON = "project.json"
PENDING_JSON = "pending.json"
SCHEDULE_JSON = "schedule.json"
MAILBOX_STATE_JSON = "mailbox_state.json"
TRACKER_FIELDNAMES = [
    "序号",
    "采购包名称",
    "供应商名称",
    "供应商简称",
    "对接人邮箱",
    "采购包类型",
    "金额不含税",
    "金额含税",
    "金额备注",
    "交期摘要",
    "已请金额",
    "待请款金额",
    "下一付款节点",
    "请款进度",
    "RFP 发出",
    "技术方案已确认",
    "商务报价已确认",
    "PO 已发起",
    "已结算",
]
PAYMENT_CSV_HEADERS = [
    "已请金额",
    "待请款金额",
    "下一付款节点",
    "请款进度",
]

# Milestone keys in chronological order; collectively represent 状态（无单独状态列）
MILESTONE_KEYS = [
    "ts_rfp_issued",
    "ts_tech_confirmed",
    "ts_quote_confirmed",
    "ts_po_issued",
    "ts_settled",
]
MILESTONE_TO_CSV = {
    "ts_rfp_issued": "RFP 发出",
    "ts_tech_confirmed": "技术方案已确认",
    "ts_quote_confirmed": "商务报价已确认",
    "ts_po_issued": "PO 已发起",
    "ts_settled": "已结算",
}
CSV_TO_MILESTONE_KEY = {v: k for k, v in MILESTONE_TO_CSV.items()}

# CLI / legacy status strings → milestone key (old 待* states mapped to completed milestone)
LEGACY_STATUS_TO_MILESTONE = {
    "待发包": None,
    "待技术确认": "ts_rfp_issued",
    "待报价确认": "ts_tech_confirmed",
    "已发PO": "ts_po_issued",
    "已结算": "ts_settled",
}

# Internal English keys <-> CSV Chinese headers
PKG_KEY_TO_CSV = {
    "seq": "序号",
    "package_name": "采购包名称",
    "supplier_name": "供应商名称",
    "supplier_short": "供应商简称",
    "contact_email": "对接人邮箱",
    "package_type": "采购包类型",
    "amount_excl_tax": "金额不含税",
    "amount_incl_tax": "金额含税",
    "amount_remark": "金额备注",
    "delivery_summary": "交期摘要",
    "claimed_amount_excl_tax": "已请金额",
    "remaining_claim_amount_excl_tax": "待请款金额",
    "next_payment_node": "下一付款节点",
    "payment_progress": "请款进度",
    "ts_rfp_issued": "RFP 发出",
    "ts_tech_confirmed": "技术方案已确认",
    "ts_quote_confirmed": "商务报价已确认",
    "ts_po_issued": "PO 已发起",
    "ts_settled": "已结算",
}
CSV_TO_PKG_KEY = {v: k for k, v in PKG_KEY_TO_CSV.items()}
# Legacy column aliases (old tracker.csv / 已取消的给人看列)
CSV_TO_PKG_KEY["总金额"] = "amount_excl_tax"
CSV_TO_PKG_KEY["状态"] = "_legacy_status"
CSV_TO_PKG_KEY["卡片路径"] = "card_path"
CSV_TO_PKG_KEY["文件夹路径"] = "folder_path"
LEGACY_TS_COLUMN_MAP = {
    "时间_待发包": "ts_rfp_issued",
    "时间_待技术确认": "ts_rfp_issued",
    "时间_待报价确认": "ts_tech_confirmed",
    "时间_已发PO": "ts_po_issued",
    "时间_已结算": "ts_settled",
}

TYPE_LABELS = {
    "供应": "纯供应",
    "供装": "供应+安装",
    "顾问": "顾问服务",
    "总包": "总包",
}

INBOX_SUBDIRS = ["raw", "attachments", "processed", "failed", "agent_mail"]


def now_iso() -> str:
    return datetime.now(CST).isoformat(timespec="seconds")


def load_json(path: Path, default: Any = None) -> Any:
    if not path.exists():
        return default
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def save_json(path: Path, data: Any) -> None:
    target = Path(path)
    if target.name == PENDING_JSON:
        raise PermissionError(
            "禁止直接写 pending.json。必须走 pending_store.add_pending / close_pending，"
            "并声明是哪个 Skill 的；进度待确认写在 schedule.json，不要混进来。"
        )
    if target.name == CARD_DATA_FILENAME:
        from write_guard import assert_card_data_write

        assert_card_data_write()
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def sanitize_segment(text: str, max_len: int = 10) -> str:
    text = re.sub(r'[\\/:*?"<>|]', "", text.strip())
    text = re.sub(r"\s+", "", text)
    return text[:max_len] if text else "未命名"


def sanitize_filename(text: str, max_len: int = 80) -> str:
    """保留空格，仅去掉文件系统非法字符。"""
    text = re.sub(r'[\\/:*?"<>|]', "", text.strip())
    return text[:max_len] if text else "project"


def tracker_basename(project: dict[str, Any]) -> str:
    """兼容旧名：返回旧版根目录 CSV 文件名（迁移用）。"""
    name = sanitize_filename(project.get("project_name") or "project")
    return f"{name}{TRACKER_CSV_LEGACY_SUFFIX}"


def tracker_xlsx_basename(project: dict[str, Any]) -> str:
    name = sanitize_filename(project.get("project_name") or "project")
    return f"{name}{TRACKER_XLSX_SUFFIX}"


def change_log_xlsx_basename(project: dict[str, Any]) -> str:
    name = sanitize_filename(project.get("project_name") or "project")
    return f"{name}{CHANGE_LOG_XLSX_SUFFIX}"


def project_archive_dir(project_root: Path, *, create: bool = False) -> Path:
    root = project_root.expanduser().resolve()
    for name in (PROJECT_ARCHIVE_DIR, *LEGACY_ARCHIVE_DIRS):
        path = root / name
        if path.is_dir():
            return path
    path = root / PROJECT_ARCHIVE_DIR
    if create:
        path.mkdir(parents=True, exist_ok=True)
    return path


def system_dir(project_root: Path, *, create: bool = False) -> Path:
    """Skill 内部目录 `.skill/`（类似 .git，Finder 默认隐藏）；兼容旧 `_系统/` 与根目录平铺。"""
    root = project_root.expanduser().resolve()
    for name in (SYSTEM_DIR, *LEGACY_SYSTEM_DIRS):
        nested = root / name
        if nested.is_dir():
            return nested
    if (root / PROJECT_JSON).is_file():
        return root
    nested = root / SYSTEM_DIR
    if create:
        nested.mkdir(parents=True, exist_ok=True)
    return nested


def system_file(project_root: Path, name: str) -> Path:
    return system_dir(project_root) / name


def package_folder_name(seq: int, package_name: str, supplier_short: str) -> str:
    """命名：{序号}_{采购包名称}_{供应商简称}，序号与包内子目录 01_ 风格一致。"""
    return (
        f"{seq:02d}_{sanitize_segment(package_name, 40)}_"
        f"{sanitize_segment(supplier_short, 20)}"
    )


def hydrate_package_paths(project_root: Path, pkg: dict[str, Any]) -> None:
    """台账不再写路径列；脚本仍靠 folder_path / card_path 找包文件夹。"""
    root = project_root.expanduser().resolve()
    pkg.setdefault("folder_path", "")
    pkg.setdefault("card_path", "")

    def _accept(folder: Path) -> None:
        resolved = folder.expanduser().resolve()
        pkg["folder_path"] = str(resolved)
        card = str(pkg.get("card_path") or "").strip()
        card_path = Path(card).expanduser() if card else resolved / CARD_FILENAME
        if card and not card_path.is_absolute():
            card_path = root / card_path
        if not card_path.is_file():
            card_path = resolved / CARD_FILENAME
        pkg["card_path"] = str(card_path)

    raw_folder = str(pkg.get("folder_path") or "").strip()
    if raw_folder:
        folder = Path(raw_folder).expanduser()
        if not folder.is_absolute():
            folder = root / folder
        if folder.suffix.lower() in {".html", ".htm", ".json"}:
            folder = folder.parent
        if folder.is_dir():
            _accept(folder)
            return

    seq_int = 0
    seq_raw = str(pkg.get("seq") or "").strip()
    if seq_raw.isdigit():
        seq_int = int(seq_raw)
    candidates: list[Path] = []
    if seq_int:
        expected = root / package_folder_name(
            seq_int,
            str(pkg.get("package_name") or ""),
            str(pkg.get("supplier_short") or ""),
        )
        if expected.is_dir():
            candidates.append(expected)
        prefix = f"{seq_int:02d}_"
        try:
            candidates.extend(
                sorted(d for d in root.iterdir() if d.is_dir() and d.name.startswith(prefix))
            )
        except OSError:
            pass
    seen: list[Path] = []
    for item in candidates:
        try:
            resolved = item.resolve()
        except OSError:
            continue
        if resolved not in seen and resolved.is_dir():
            seen.append(resolved)
    if seen:
        _accept(seen[0])


def render_project_background_md(project: dict[str, Any]) -> str:
    lines = [
        f"# {project.get('project_name', '未命名项目')}",
        "",
        f"- **项目编号**：{project.get('project_code', '—')}",
        f"- **项目类型**：{project.get('project_type', '—')}",
        f"- **项目地址**：{project.get('address', '—')}",
        f"- **甲方（业主）**：{project.get('client_name', '—')}",
        f"- **项目经理**：{project.get('project_manager', '—')}",
        f"- **币种 / 含税约定**：{project.get('currency_tax', '—')}",
        f"- **初始化时间**：{project.get('initialized_at', '—')}",
    ]
    mailbox = project.get("mailbox") or {}
    if mailbox.get("address"):
        lines.append(f"- **Skill 专用邮箱**：{mailbox['address']}")
    lines.extend(
        [
            "",
            "## 进度计划",
            "",
            "完整进度计划文件见 `项目档案/00_进度计划/`。",
            "",
            "> 本文件由 Skill 自动生成，供 PM 日常查阅。机器可读配置在 `.skill/project.json`。",
        ]
    )
    tracker_updated = project.get("tracker_last_updated")
    tracker_file = project.get("tracker_filename")
    if tracker_updated or tracker_file:
        lines.extend(["", "## Tracker", ""])
        if tracker_file:
            lines.append(f"- **Tracker 文件**：`{tracker_file}`")
        if tracker_updated:
            lines.append(f"- **Tracker 最后更新**：{tracker_updated}")
    lines.append("")
    return "\n".join(lines) + "\n"


def write_project_background_md(project_root: Path, project: dict[str, Any] | None = None) -> Path:
    root = project_root.expanduser().resolve()
    if project is None:
        project = load_json(system_file(root, PROJECT_JSON), {})
    path = root / BACKGROUND_MD
    path.write_text(render_project_background_md(project or {}), encoding="utf-8")
    return path


def subdirs_for_type(package_type: str) -> list[str]:
    dirs = list(PACKAGE_SUBDIRS_ALL)
    if package_type == "总包":
        dirs.append(PACKAGE_SUBDIR_GC)
    if package_type == "顾问":
        dirs.append(PACKAGE_SUBDIR_CONSULTANT)
    return dirs


def default_card_data(
    *,
    package_name: str,
    supplier_name: str,
    package_type: str,
    project_name: str,
) -> dict[str, Any]:
    return {
        "package_name": package_name,
        "supplier_name": supplier_name,
        "package_type": package_type,
        "package_type_label": TYPE_LABELS.get(package_type, package_type),
        "project_name": project_name,
        "contact_email": "",
        "total_amount": "",
        "scope_summary": "",
        "delivery_schedule": "",
        "payment_nodes": [],
        "change_items": [],
        "acceptance_items": [],
        "constraint_clauses": "",
        "data_sources": "",
    }


def deep_merge(base: dict[str, Any], patch: dict[str, Any]) -> dict[str, Any]:
    result = dict(base)
    for key, value in patch.items():
        if (
            key in result
            and isinstance(result[key], dict)
            and isinstance(value, dict)
        ):
            result[key] = deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def format_milestone_date(value: str) -> str:
    """Normalize milestone values to YYYY-MM-DD (no time-of-day)."""
    if not (value or "").strip():
        return ""
    v = value.strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}$", v):
        return v
    try:
        dt = datetime.fromisoformat(v.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=CST)
        return dt.astimezone(CST).strftime("%Y-%m-%d")
    except ValueError:
        return v[:10] if len(v) >= 10 else v


def split_amount_and_remark(text: str) -> tuple[str, str]:
    """Return (numeric amount, remark). Amount is plain decimal string."""
    if not (text or "").strip():
        return "", ""
    raw = text.strip()
    remarks: list[str] = []
    for match in re.finditer(r"[（(]([^）)]+)[）)]", raw):
        part = match.group(1).strip()
        if part:
            remarks.append(part)
    without_paren = re.sub(r"[（(][^）)]+[）)]", "", raw).strip()
    num_match = re.search(r"([\d,]+(?:\.\d+)?)", without_paren.replace(" ", ""))
    if not num_match:
        return "", "; ".join(remarks) if remarks else raw
    num = float(num_match.group(1).replace(",", ""))
    amount = f"{num:.2f}"
    residual = without_paren
    residual = re.sub(r"[¥￥$]\s*", "", residual)
    residual = re.sub(r"[\d,]+(?:\.\d+)?", "", residual)
    residual = re.sub(
        r"\b(CNY|RMB|USD|含税|不含税|excl\.?\s*VAT|incl\.?\s*VAT)\b",
        "",
        residual,
        flags=re.I,
    )
    residual = re.sub(r"\s+", " ", residual).strip(" ;，,")
    if residual:
        remarks.append(residual)
    remark = "; ".join(dict.fromkeys(r for r in remarks if r))
    return amount, remark


def infer_package_seq(pkg: dict[str, Any], fallback: int) -> str:
    folder = str(pkg.get("folder_path") or "")
    name = Path(folder).name if folder else ""
    match = re.match(r"^(\d{2})_", name)
    if match:
        return match.group(1)
    existing = str(pkg.get("seq") or "").strip()
    if existing.isdigit():
        return f"{int(existing):02d}"
    return f"{fallback:02d}"


def normalize_pkg_amounts(pkg: dict[str, Any]) -> None:
    excl, excl_remark = split_amount_and_remark(str(pkg.get("amount_excl_tax") or ""))
    incl, incl_remark = split_amount_and_remark(str(pkg.get("amount_incl_tax") or ""))
    pkg["amount_excl_tax"] = excl
    pkg["amount_incl_tax"] = incl
    remark_parts: list[str] = []
    for part in [pkg.get("amount_remark"), excl_remark, incl_remark]:
        if not part:
            continue
        for piece in re.split(r"[;；]", str(part)):
            piece = piece.strip()
            if piece and piece not in remark_parts:
                remark_parts.append(piece)
    pkg["amount_remark"] = "; ".join(remark_parts)


def normalize_pkg_for_save(pkg: dict[str, Any], index: int) -> None:
    normalize_pkg_amounts(pkg)
    pkg["seq"] = infer_package_seq(pkg, index)
    for key in MILESTONE_KEYS:
        pkg[key] = format_milestone_date(str(pkg.get(key) or ""))


def pkg_to_csv_row(pkg: dict[str, Any]) -> dict[str, str]:
    row = {csv_key: str(pkg.get(eng_key, "") or "") for eng_key, csv_key in PKG_KEY_TO_CSV.items()}
    for col in MILESTONE_TO_CSV.values():
        if row.get(col):
            row[col] = format_milestone_date(row[col])
    return row


def _milestone_column_range() -> tuple[int, int]:
    first = TRACKER_FIELDNAMES.index(MILESTONE_TO_CSV["ts_rfp_issued"])
    last = TRACKER_FIELDNAMES.index(MILESTONE_TO_CSV["ts_settled"])
    return first, last


def _payment_column_range() -> tuple[int, int]:
    first = TRACKER_FIELDNAMES.index("已请金额")
    last = TRACKER_FIELDNAMES.index("请款进度")
    return first, last


def build_tracker_group_row() -> list[str]:
    """CSV / Excel 分组行：请款章节标「请款」，里程碑标「状态」。"""
    row = [""] * len(TRACKER_FIELDNAMES)
    pay_first, _ = _payment_column_range()
    ms_first, _ = _milestone_column_range()
    row[pay_first] = "请款"
    row[ms_first] = "状态"
    return row


def _is_tracker_group_header_line(line: str) -> bool:
    parsed = next(csv.reader([line]), [])
    if not parsed:
        return False
    if (parsed[0] or "").strip() == "序号":
        return False
    ms_first, _ = _milestone_column_range()
    if ms_first < len(parsed) and (parsed[ms_first] or "").strip() == "状态":
        return True
    return "状态" in parsed and "采购包名称" not in parsed


def csv_row_to_pkg(row: dict[str, str]) -> dict[str, str]:
    pkg: dict[str, str] = {}
    for k, v in row.items():
        if k in CSV_TO_PKG_KEY:
            pkg[CSV_TO_PKG_KEY[k]] = v or ""
        elif k in LEGACY_TS_COLUMN_MAP:
            key = LEGACY_TS_COLUMN_MAP[k]
            if v and not (pkg.get(key) or "").strip():
                pkg[key] = v
    legacy_status = (pkg.pop("_legacy_status", "") or row.get("状态", "") or "").strip()
    if legacy_status and not any((pkg.get(k) or "").strip() for k in MILESTONE_KEYS):
        milestone = LEGACY_STATUS_TO_MILESTONE.get(legacy_status)
        if milestone:
            pkg.setdefault(milestone, "")
    if (pkg.get("ts_po_issued") or "").strip() and not (pkg.get("ts_quote_confirmed") or "").strip():
        pkg["ts_quote_confirmed"] = pkg["ts_po_issued"]
    return pkg


def resolve_tracker_path(project_root: Path) -> Path | None:
    """定位 Tracker 机读 CSV（优先 `.skill/status_tracker.csv`）；不存在时返回将要写入的目标路径。"""
    root = project_root.expanduser().resolve()
    project = load_json(system_file(root, PROJECT_JSON), {}) or {}
    skill = system_dir(root, create=False)
    preferred = skill / TRACKER_CSV_NAME

    candidates: list[Path] = [preferred]
    if project.get("tracker_filename"):
        # 旧 project.json 可能记根目录相对名，或 .skill/ 相对名
        tf = str(project["tracker_filename"])
        candidates.append(root / tf)
        candidates.append(skill / Path(tf).name)
    if project.get("project_name"):
        candidates.append(root / tracker_basename(project))
    candidates.append(root / TRACKER_CSV_LEGACY)
    candidates.extend(sorted(root.glob(f"*{TRACKER_CSV_LEGACY_SUFFIX}")))
    candidates.extend(sorted(skill.glob(f"*{TRACKER_CSV_LEGACY_SUFFIX}")))

    seen: set[Path] = set()
    for path in candidates:
        path = path.resolve()
        if path in seen:
            continue
        seen.add(path)
        if path.is_file():
            return path

    return preferred if project.get("project_name") or skill.is_dir() else root / TRACKER_CSV_LEGACY


def resolve_tracker_xlsx_path(project_root: Path) -> Path:
    """给人看的 Status Tracker Excel（项目根）。"""
    root = project_root.expanduser().resolve()
    project = load_json(system_file(root, PROJECT_JSON), {}) or {}
    if project.get("tracker_xlsx"):
        p = root / str(project["tracker_xlsx"])
        if p.is_file() or project.get("project_name"):
            # 旧名仍可读；新写用标准名
            pass
    if project.get("project_name"):
        return root / tracker_xlsx_basename(project)
    return root / f"project{TRACKER_XLSX_SUFFIX}"


def _read_tracker_csv_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8-sig")
    lines = [line for line in raw.splitlines(keepends=True) if not line.lstrip().startswith("#")]
    if lines and _is_tracker_group_header_line(lines[0].rstrip("\r\n")):
        lines = lines[1:]
    return "".join(lines)


def load_tracker_packages(project_root: Path) -> list[dict[str, str]]:
    path = resolve_tracker_path(project_root)
    if not path or not path.exists():
        return []
    text = _read_tracker_csv_text(path)
    reader = csv.DictReader(io.StringIO(text))
    packages = [csv_row_to_pkg(row) for row in reader if any(row.values())]
    for idx, pkg in enumerate(packages, start=1):
        for key in PKG_KEY_TO_CSV:
            pkg.setdefault(key, "")
        if pkg.get("total_amount") and not pkg.get("amount_excl_tax"):
            pkg["amount_excl_tax"] = pkg["total_amount"]
        normalize_pkg_amounts(pkg)
        if not (pkg.get("seq") or "").strip():
            pkg["seq"] = infer_package_seq(pkg, idx)
        hydrate_package_paths(project_root, pkg)
        for key in MILESTONE_KEYS:
            if pkg.get(key):
                pkg[key] = format_milestone_date(str(pkg[key]))
    return packages


def resolve_change_log_path(project_root: Path) -> Path | None:
    """机读 Change Log CSV（优先 `.skill/change_log.csv`）。无项目壳则 None。"""
    root = project_root.expanduser().resolve()
    project = load_json(system_file(root, PROJECT_JSON), {}) or {}
    skill = system_dir(root, create=False)
    preferred = skill / CHANGE_LOG_CSV_NAME

    candidates: list[Path] = [preferred]
    name = str(project.get("project_name") or "").strip()
    if name:
        safe = sanitize_filename(name)
        candidates.append(root / f"{safe}{CHANGE_LOG_CSV_LEGACY_SUFFIX}")
        candidates.append(skill / f"{safe}{CHANGE_LOG_CSV_LEGACY_SUFFIX}")
    candidates.extend(sorted(root.glob(f"*{CHANGE_LOG_CSV_LEGACY_SUFFIX}")))
    candidates.extend(sorted(skill.glob(f"*{CHANGE_LOG_CSV_LEGACY_SUFFIX}")))

    seen: set[Path] = set()
    for path in candidates:
        path = path.resolve()
        if path in seen:
            continue
        seen.add(path)
        if path.is_file():
            return path

    if not name and not skill.is_dir() and not (root / PROJECT_JSON).is_file() and not system_file(root, PROJECT_JSON).is_file():
        return None
    return preferred


def resolve_change_log_xlsx_path(project_root: Path) -> Path | None:
    """给人看的 Change Log Excel（项目根）。"""
    root = project_root.expanduser().resolve()
    project = load_json(system_file(root, PROJECT_JSON), {}) or {}
    name = str(project.get("project_name") or "").strip()
    if not name:
        return None
    return root / change_log_xlsx_basename(project)


def _change_log_package_key(text: str) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", (text or "").lower())


def _change_log_package_key_variants(name: str) -> set[str]:
    key = _change_log_package_key(name)
    out = {key} if key else set()
    if not key:
        return out
    if key.endswith("ies") and len(key) > 4:
        out.add(key[:-3] + "y")
    elif key.endswith("s") and len(key) > 3 and not key.endswith(("ss", "us", "is")):
        out.add(key[:-1])
    else:
        out.add(key + "s")
    return {k for k in out if k}


def _change_log_packages_match(a: str, b: str) -> bool:
    if not a or not b:
        return False
    if a.strip() == b.strip():
        return True
    return bool(_change_log_package_key_variants(a) & _change_log_package_key_variants(b))


def _change_log_row_cancelled(row: dict[str, str]) -> bool:
    if str(row.get("当前状态") or "").strip() == "已取消":
        return True
    return bool(str(row.get("已取消") or "").strip())


def load_change_log_rows(project_root: Path) -> list[dict[str, str]]:
    """读 Change Log 数据行（跳过 # 元数据行）。文件不存在 → []。"""
    path = resolve_change_log_path(project_root)
    if not path or not path.is_file():
        return []
    lines = [
        ln for ln in path.read_text(encoding="utf-8-sig").splitlines() if ln.strip() and not ln.lstrip().startswith("#")
    ]
    if not lines:
        return []
    reader = csv.DictReader(io.StringIO("\n".join(lines)))
    return [dict(row) for row in reader if any((v or "").strip() for v in row.values())]


def confirmed_changes_for_package(project_root: Path, package_name: str) -> list[dict[str, Any]]:
    """该包已确认变更（有最终确认金额、非已取消）。供现金流 / 请款页只读。

    每项：number / content / confirmed_amount（原文字符串）/ amount_excl_tax（float|None）。
    """
    out: list[dict[str, Any]] = []
    for row in load_change_log_rows(project_root):
        if _change_log_row_cancelled(row):
            continue
        if not _change_log_packages_match(str(row.get("采购包名称") or ""), package_name):
            continue
        raw = str(row.get("最终确认金额") or "").strip()
        if not raw:
            continue
        parsed = parse_amount_with_tax(raw)
        amount: float | None = None
        for key in ("amount_excl_tax", "amount", "value"):
            v = parsed.get(key)
            if v is None or v == "":
                continue
            try:
                amount = float(str(v).replace(",", ""))
                break
            except (TypeError, ValueError):
                pass
        if amount is None:
            m = re.search(r"([\d,]+(?:\.\d+)?)", raw.replace("¥", "").replace("￥", "").replace(" ", ""))
            if m:
                try:
                    amount = float(m.group(1).replace(",", ""))
                except ValueError:
                    amount = None
        out.append(
            {
                "number": str(row.get("变更序号") or "").strip(),
                "content": str(row.get("变更内容") or "").strip(),
                "confirmed_amount": raw,
                "amount_excl_tax": amount,
            }
        )
    return out


def confirmed_changes_total_excl_tax(project_root: Path, package_name: str) -> float:
    total = 0.0
    for item in confirmed_changes_for_package(project_root, package_name):
        amt = item.get("amount_excl_tax")
        if isinstance(amt, (int, float)):
            total += float(amt)
    return total


def _cleanup_legacy_tracker_files(project_root: Path, keep_csv: Path, keep_xlsx: Path) -> None:
    root = project_root.expanduser().resolve()
    skill = system_dir(root, create=False)
    keep_csv = keep_csv.resolve()
    keep_xlsx = keep_xlsx.resolve()
    csv_globs = [
        root / TRACKER_CSV_LEGACY,
        *root.glob(f"*{TRACKER_CSV_LEGACY_SUFFIX}"),
        *skill.glob(f"*{TRACKER_CSV_LEGACY_SUFFIX}"),
        skill / TRACKER_CSV_NAME,
    ]
    for path in csv_globs:
        path = path.resolve()
        if path != keep_csv and path.is_file():
            path.unlink()
    for path in [
        *root.glob("*_procurement_packages_status_tracker.xlsx"),
        *root.glob(f"*{TRACKER_XLSX_SUFFIX}"),
    ]:
        path = path.resolve()
        if path != keep_xlsx and path.is_file():
            path.unlink()


def save_tracker_packages(project_root: Path, packages: list[dict[str, Any]]) -> Path:
    root = project_root.expanduser().resolve()
    from write_guard import assert_tracker_write

    assert_tracker_write(root, packages)
    project = load_json(system_file(root, PROJECT_JSON), {}) or {}
    skill = system_dir(root, create=True)
    path = skill / TRACKER_CSV_NAME
    xlsx_path = resolve_tracker_xlsx_path(root)
    updated_at = now_iso()

    path.parent.mkdir(parents=True, exist_ok=True)
    prepared: list[dict[str, Any]] = []
    for idx, pkg in enumerate(packages, start=1):
        for key in PKG_KEY_TO_CSV:
            pkg.setdefault(key, "")
        normalize_pkg_for_save(pkg, idx)
        prepared.append(pkg)

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        f.write(f"# 最后更新,{updated_at}\n")
        writer = csv.writer(f)
        writer.writerow(build_tracker_group_row())
        writer.writerow(TRACKER_FIELDNAMES)
        for pkg in prepared:
            row = pkg_to_csv_row(pkg)
            writer.writerow([row.get(col, "") for col in TRACKER_FIELDNAMES])

    from tracker_xlsx import save_tracker_xlsx

    save_tracker_xlsx(
        xlsx_path,
        prepared,
        updated_at,
        project_name=str(project.get("project_name") or ""),
    )

    project["tracker_filename"] = f"{SYSTEM_DIR}/{TRACKER_CSV_NAME}"
    project["tracker_xlsx"] = xlsx_path.name
    project["tracker_last_updated"] = updated_at
    save_json(system_file(root, PROJECT_JSON), project)
    write_project_background_md(root, project)
    _cleanup_legacy_tracker_files(root, path, xlsx_path)
    return path


def resolve_milestone_key(milestone: str) -> str | None:
    """Accept internal key, CSV column name, or legacy status label."""
    if milestone in MILESTONE_TO_CSV:
        return milestone
    if milestone in CSV_TO_MILESTONE_KEY:
        return CSV_TO_MILESTONE_KEY[milestone]
    return LEGACY_STATUS_TO_MILESTONE.get(milestone)


def stamp_milestone(pkg: dict[str, Any], milestone: str, when: str | None = None) -> None:
    """Record completion datetime for a milestone (first write wins)."""
    key = resolve_milestone_key(milestone)
    if not key:
        return
    when = milestone_date_value(when)
    if not (pkg.get(key) or "").strip():
        pkg[key] = when
    # PO 发起隐含商务报价已确认
    if key == "ts_po_issued" and not (pkg.get("ts_quote_confirmed") or "").strip():
        pkg["ts_quote_confirmed"] = when


def milestone_date_value(when: str | None = None) -> str:
    return format_milestone_date(when or now_iso())


def latest_milestone_label(pkg: dict[str, Any]) -> str:
    """Return the furthest completed milestone CSV label, or empty if none."""
    for key in reversed(MILESTONE_KEYS):
        if (pkg.get(key) or "").strip():
            return MILESTONE_TO_CSV[key]
    return ""


def milestone_reached(pkg: dict[str, Any], milestone: str) -> bool:
    key = resolve_milestone_key(milestone)
    return bool(key and (pkg.get(key) or "").strip())


def format_currency_amount(amount: float | str, currency: str = "CNY") -> str:
    """Unified money display: ¥12,345.67 (CNY/RMB) or USD 12,345.67."""
    if isinstance(amount, str):
        raw = amount.strip()
        if not raw:
            return ""
        m = re.search(r"([\d,]+(?:\.\d+)?)", raw.replace(" ", ""))
        if not m:
            return raw
        num = float(m.group(1).replace(",", ""))
        cur = currency
        if re.search(r"USD|\$", raw, re.I):
            cur = "USD"
        elif re.search(r"RMB|CNY|¥|￥", raw, re.I):
            cur = "CNY"
    else:
        num = float(amount)
        cur = currency
    formatted = f"{num:,.2f}"
    if cur.upper() in ("CNY", "RMB", "¥", "￥"):
        return f"¥{formatted}"
    return f"{cur.upper()} {formatted}"


def parse_amount_with_tax(text: str) -> dict[str, Any]:
    """
    Parse amount + tax hint from free text.
    Returns amount_excl_tax, amount_incl_tax (formatted), tax_rate, tax_uncertain, raw.
    """
    out: dict[str, Any] = {
        "amount_excl_tax": "",
        "amount_incl_tax": "",
        "tax_rate": None,
        "tax_uncertain": False,
        "raw": (text or "").strip(),
    }
    if not text or not text.strip():
        return out
    raw = text.strip()
    m_num = re.search(r"(RMB|CNY|USD|¥|￥)?\s*([\d,]+(?:\.\d+)?)", raw, re.I)
    if not m_num:
        out["tax_uncertain"] = True
        return out
    cur = (m_num.group(1) or "CNY").upper().replace("￥", "¥")
    if cur in ("¥", "￥", "RMB"):
        cur = "CNY"
    num = float(m_num.group(2).replace(",", ""))

    excl = bool(re.search(r"excl\.?\s*|不含\s*税|未税|不含增值税|before\s*VAT|Amount\s*before", raw, re.I))
    incl = bool(re.search(r"incl\.?\s*|含\s*税|含增值税|Include\.?\s*VAT|Amount\s*\(?\s*Include", raw, re.I))
    m_rate = re.search(
        r"(\d+(?:\.\d+)?)\s*%\s*(?:VAT|vat|tax|Tax|TAX|税)",
        raw,
    )
    rate = float(m_rate.group(1)) / 100.0 if m_rate else None

    if excl and rate is not None:
        out["amount_excl_tax"] = format_currency_amount(num, cur)
        out["amount_incl_tax"] = format_currency_amount(num * (1 + rate), cur)
        out["tax_rate"] = rate
    elif incl and rate is not None:
        out["amount_incl_tax"] = format_currency_amount(num, cur)
        out["amount_excl_tax"] = format_currency_amount(num / (1 + rate), cur)
        out["tax_rate"] = rate
    elif excl:
        out["amount_excl_tax"] = format_currency_amount(num, cur)
        out["tax_uncertain"] = True  # need rate or incl amount
    elif incl:
        out["amount_incl_tax"] = format_currency_amount(num, cur)
        out["tax_uncertain"] = True
    else:
        # Ambiguous — keep as excl placeholder and flag
        out["amount_excl_tax"] = format_currency_amount(num, cur)
        out["tax_uncertain"] = True
    return out


def _sheet_rows_as_text_cells(path: Path) -> list[list[Any]]:
    """Load .xls / .xlsx into list of rows (best-effort)."""
    ext = path.suffix.lower()
    rows: list[list[Any]] = []
    if ext == ".xls":
        try:
            import xlrd  # type: ignore
        except ImportError:
            return rows
        book = xlrd.open_workbook(str(path))
        for sheet in book.sheets():
            for r in range(sheet.nrows):
                rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
        return rows
    if ext in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError:
            return rows
        wb = load_workbook(str(path), data_only=True, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
        return rows
    return rows


def extract_amounts_from_quote_file(path: Path) -> dict[str, Any]:
    """
    Pull 不含税 / 含税 from quotation spreadsheets.

    Supports label-in-row sheets (e.g. SecuALL Amount before VAT) and
    header-column sheets (e.g. Global Total (RMB) w/o Tax | Tax(9%) | With Tax).
    """
    out: dict[str, Any] = {
        "amount_excl_tax": "",
        "amount_incl_tax": "",
        "tax_rate": None,
        "tax_uncertain": False,
        "source": "",
    }
    path = Path(path)
    if not path.exists() or path.suffix.lower() not in {".xls", ".xlsx", ".xlsm", ".csv"}:
        return out

    excl_val: float | None = None
    incl_val: float | None = None
    vat_val: float | None = None

    # w/o tax must be checked before "with tax"; avoid matching w/o as with
    _EXCL_LAB = (
        r"amount\s*before\s*vat|before\s*vat|不含\s*税|未税|金额不含税|"
        r"w\s*/\s*o\s*tax|without\s*tax|excl\.?\s*(?:tax|vat)|"
        r"global\s*total[^\n|]{0,40}w\s*/\s*o"
    )
    _INCL_LAB = (
        r"amount\s*\(?\s*include\.?\s*vat|include\.?\s*vat|含\s*税|金额含税|"
        r"with\s*tax|incl\.?\s*(?:tax|vat)|"
        r"global\s*total[^\n|]{0,40}with\s*tax"
    )
    _VAT_LAB = r"(?:^|\W)(?:vat|tax)\s*\(?\s*\d|(?:^|\W)tax\s*（|增值税|(?:^|\|)tax(?:\||$)"

    def _num(cell: Any) -> float | None:
        if cell is None or cell == "":
            return None
        if isinstance(cell, (int, float)):
            return float(cell)
        text = str(cell).strip()
        # "Tax（9%）" / "Freight (%)" — rate labels, not money
        if re.search(r"[A-Za-z\u4e00-\u9fff]", text) and re.search(
            r"tax|vat|%|％|税|freight|duty|markup", text, re.I
        ):
            return None
        compact = text.replace(" ", "").replace(",", "")
        m = re.search(r"(\d+(?:\.\d+)?)", compact)
        if not m:
            return None
        try:
            return float(m.group(1))
        except ValueError:
            return None

    def _is_excl_label(text: str) -> bool:
        return bool(re.search(_EXCL_LAB, text, re.I))

    def _is_incl_label(text: str) -> bool:
        if re.search(r"w\s*/\s*o\s*tax|without\s*tax|before\s*vat|不含", text, re.I):
            return False
        return bool(re.search(_INCL_LAB, text, re.I))

    def _is_vat_label(text: str) -> bool:
        if re.search(r"before|include|amount|w\s*/\s*o|with\s*tax|不含|含税", text, re.I):
            return False
        return bool(re.search(_VAT_LAB, text, re.I)) or bool(
            re.search(r"(?:^|\W)vat(?:\W|$)", text, re.I)
        )

    def _at(row: list[Any], idx: int | None) -> float | None:
        if idx is None or idx >= len(row):
            return None
        return _num(row[idx])

    rows = _sheet_rows_as_text_cells(path)
    if path.suffix.lower() == ".csv" and not rows:
        try:
            import csv as _csv

            with path.open(encoding="utf-8-sig", newline="") as f:
                rows = [list(r) for r in _csv.reader(f)]
        except OSError:
            rows = []

    # Pass 1: header columns → pick grand-total row (max excl, or last sparse total row)
    for hi, row in enumerate(rows):
        headers = [("" if c is None else str(c)).strip() for c in row]
        if not any(headers):
            continue
        excl_i = next((i for i, h in enumerate(headers) if _is_excl_label(h)), None)
        incl_i = next((i for i, h in enumerate(headers) if _is_incl_label(h)), None)
        vat_i = next((i for i, h in enumerate(headers) if _is_vat_label(h)), None)
        if excl_i is None and incl_i is None:
            continue

        scored: list[tuple[float, float | None, float | None, float | None]] = []
        for data in rows[hi + 1 :]:
            e, i_amt, v = _at(data, excl_i), _at(data, incl_i), _at(data, vat_i)
            if e is None and i_amt is None:
                continue
            # Prefer rows that look like totals (little else filled) via boost
            filled = sum(1 for c in data if c not in (None, ""))
            score = float(e or i_amt or 0)
            if filled <= 4:
                score += 1e12  # total-only row beats line items
            scored.append((score, e, i_amt, v))
        if scored:
            scored.sort(key=lambda t: t[0])
            _, e, i_amt, v = scored[-1]
            if e is not None:
                excl_val = e
            if i_amt is not None:
                incl_val = i_amt
            if v is not None:
                vat_val = v
            break

    # Pass 2: label + amount in the same row (legacy SecuALL-style)
    if excl_val is None and incl_val is None:
        for row in rows:
            cells = [("" if c is None else str(c)).strip() for c in row]
            joined = " | ".join(cells)
            nums = [n for n in (_num(c) for c in row) if n is not None]
            if not nums:
                continue
            if _is_excl_label(joined):
                for n in reversed(nums):
                    if n > 0:
                        excl_val = n
                        break
            if _is_incl_label(joined):
                for n in reversed(nums):
                    if n > 0:
                        incl_val = n
                        break
            if _is_vat_label(joined):
                for n in reversed(nums):
                    if n > 0:
                        vat_val = n
                        break
            if len(cells) >= 2:
                lab0 = cells[0]
                n_last = _num(row[-1]) if row else None
                if n_last and _is_excl_label(lab0):
                    excl_val = n_last
                if n_last and _is_incl_label(lab0):
                    incl_val = n_last

    if excl_val is not None:
        out["amount_excl_tax"] = format_currency_amount(excl_val)
    if incl_val is not None:
        out["amount_incl_tax"] = format_currency_amount(incl_val)
    if excl_val is not None and incl_val is not None and excl_val > 0:
        out["tax_rate"] = (incl_val - excl_val) / excl_val
        out["tax_uncertain"] = False
    elif excl_val is not None and vat_val is not None:
        out["amount_incl_tax"] = format_currency_amount(excl_val + vat_val)
        out["tax_rate"] = vat_val / excl_val if excl_val else None
        out["tax_uncertain"] = False
    elif excl_val is not None or incl_val is not None:
        out["tax_uncertain"] = not (excl_val is not None and incl_val is not None)

    if out["amount_excl_tax"] or out["amount_incl_tax"]:
        out["source"] = str(path)
    return out


def is_quotation_attachment(path: Path) -> bool:
    """Keep quote/contract docs under 01_合同报价; never store email signature images there."""
    name = path.name
    ext = path.suffix.lower()
    quote_exts = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".csv", ".zip", ".rar", ".7z"}
    image_exts = {".gif", ".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
    if ext in image_exts:
        return False
    if ext in quote_exts:
        return True
    # Other extensions only if name clearly quote/contract
    return bool(re.search(r"quot|报价|boq|contract|合同|invoice|请款|bill|price", name, re.I))


def merge_contact_emails(*groups: str) -> str:
    seen: list[str] = []
    for g in groups:
        if not g:
            continue
        for part in re.split(r"[;,\s]+", g):
            addr = part.strip().lower()
            if "@" in addr and addr not in seen:
                seen.append(addr)
    return "; ".join(seen)


def find_package(packages: list[dict[str, Any]], package_name: str) -> dict[str, Any] | None:
    for pkg in packages:
        if pkg.get("package_name") == package_name:
            return pkg
    return None


DROP_MAIL_DIRNAME = "待处理邮件"


def drop_mail_dir(project_root: Path, *, create: bool = False) -> Path:
    """对人可见的待处理邮件夹（项目根下）。不是 .skill/inbox。"""
    path = project_root.expanduser().resolve() / DROP_MAIL_DIRNAME
    if create:
        path.mkdir(parents=True, exist_ok=True)
    return path


def project_json_path(project_root: Path) -> Path:
    return project_root.expanduser().resolve() / ".skill" / PROJECT_JSON


def ensure_inbox_dirs(project_root: Path) -> None:
    inbox = system_file(project_root, "inbox")
    for sub in INBOX_SUBDIRS:
        (inbox / sub).mkdir(parents=True, exist_ok=True)


# Heuristic 请款附件 — overridden when node has explicit payment_attachments
_PAYMENT_ATTACHMENT_RULES: list[tuple[str, str]] = [
    (
        r"预付款|advance|down\s*payment|deposit|开工|mobilization",
        "请款单；PO/合同复印件；发票（如适用）；预付款保函（如合同要求）",
    ),
    (
        r"到货|delivery|material|设备到场|材料|shipping",
        "请款单；到货签认单；送货单/装箱单；发票",
    ),
    (
        r"安装|installation|调试|commissioning|联调",
        "请款单；安装/调试完成确认单；测试记录；发票",
    ),
    (
        r"验收|acceptance|handover|竣工|到货验收|完工",
        "请款单；验收签认单；测试/调试报告；发票",
    ),
    (
        r"进度|progress|月度|施工进度|工程量",
        "请款单；经签认的进度款申请/工程量清单；月度进度报告；发票",
    ),
    (
        r"尾款|质保|retention|final|warranty|结算|缺陷责任|服务完成|service\s*complet",
        "请款单；最终验收/质保期满证明；发票；保函释放申请（如适用）",
    ),
]

_DEFAULT_PAYMENT_ATTACHMENTS = "请款单；相关签认单或支持性文件；发票"

_ISO_DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}")

_PAYMENT_SCHEDULE_RULES: list[tuple[str, list[str], str]] = [
    (r"预付款|advance|down\s*payment|deposit", ["开工", "start", "commencement"], "before_start"),
    (r"到货|delivery|material|设备到场", ["到货", "delivery", "material"], "delivery"),
    (r"安装|install", ["安装", "installation", "install"], "range"),
    (r"调试|commission|联调", ["调试", "commission", "联调"], "range"),
    (
        r"验收|acceptance|handover|完工",
        ["验收", "acceptance", "移交", "handover"],
        "before_handover",
    ),
    (
        r"尾款|final|retention|质保|warranty|结算|服务完成|service\s*complet",
        ["竣工", "完工", "移交", "handover", "completion", "final", "服务完成"],
        "at_or_before",
    ),
]

# 付款语义 → critical_nodes.node_code（优先于 milestones 关键词）
_PAYMENT_TO_NODE_CODES: list[tuple[str, list[str]]] = [
    (r"预付款|advance|down\s*payment|deposit", ["po_confirm", "project_start"]),
    (r"到货|delivery|material|设备到场", ["delivery"]),
    (r"安装|install", ["install_complete", "commissioning"]),
    (r"调试|commission|联调", ["commissioning", "install_complete"]),
    (
        r"验收|acceptance|handover|完工",
        ["final_accept", "prelim_accept", "deliverable_accept", "project_handover"],
    ),
    (
        r"尾款|final|retention|质保|warranty|结算|服务完成|service\s*complet",
        [
            "warranty_start",
            "final_accept",
            "project_handover",
            "project_completion",
            "deliverable_accept",
        ],
    ),
]


def load_critical_nodes(project_root: Path) -> list[dict[str, Any]]:
    path = project_root.expanduser().resolve() / SYSTEM_DIR / SCHEDULE_JSON
    data = load_json(path, {})
    nodes = data.get("critical_nodes") if isinstance(data, dict) else None
    return nodes if isinstance(nodes, list) else []


def _pick_critical_node_end(
    critical_nodes: list[dict[str, Any]],
    *,
    package_name: str,
    node_codes: list[str],
) -> str:
    pkg = (package_name or "").strip().lower()
    by_code: dict[str, dict[str, Any]] = {}
    for n in critical_nodes:
        if not isinstance(n, dict):
            continue
        pname = str(n.get("package_name") or "").strip().lower()
        if pkg and pname != pkg:
            continue
        code = str(n.get("node_code") or "").strip()
        if code and code not in by_code:
            by_code[code] = n
    for code in node_codes:
        n = by_code.get(code)
        if not n:
            continue
        end = str(n.get("window_end") or n.get("window_start") or "").strip()
        if end:
            return end
    # 验收类可回退总包客户验收 / 移交
    if any(
        c
        in {
            "final_accept",
            "prelim_accept",
            "project_handover",
            "deliverable_accept",
            "warranty_start",
        }
        for c in node_codes
    ):
        for n in critical_nodes:
            if not isinstance(n, dict):
                continue
            if str(n.get("package_name") or "") not in {"总包", "_project"}:
                continue
            code = str(n.get("node_code") or "")
            name = str(n.get("name") or "") + str(n.get("name_zh") or "")
            if code == "gc_milestone" and not re.search(
                r"客户验收|final acceptance|移交|handover", name, re.I
            ):
                continue
            if code not in {
                "project_handover",
                "project_completion",
                "gc_milestone",
            }:
                continue
            end = str(n.get("window_end") or "").strip()
            if end:
                return end
    return ""


def infer_window_from_critical_nodes(
    node: dict[str, Any],
    *,
    critical_nodes: list[dict[str, Any]],
    package_name: str = "",
) -> str:
    if not critical_nodes:
        return ""
    haystack = " ".join(
        str(node.get(key) or "") for key in ("name", "condition", "milestone")
    )
    for pattern, codes in _PAYMENT_TO_NODE_CODES:
        if not re.search(pattern, haystack, re.I):
            continue
        end = _pick_critical_node_end(
            critical_nodes, package_name=package_name, node_codes=codes
        )
        if end:
            return f"≤ {end}"
    return ""


def resolve_project_root_from_package(package_dir: Path) -> Path | None:
    current = package_dir.expanduser().resolve()
    if current.is_file():
        current = current.parent
    for _ in range(6):
        skill_project = current / SYSTEM_DIR / PROJECT_JSON
        tracker = resolve_tracker_path(current)
        if skill_project.is_file() or (tracker and tracker.is_file()):
            return current
        if current.parent == current:
            break
        current = current.parent
    return None


def load_project_schedule(project_root: Path) -> list[dict[str, str]]:
    from schedule_csv import load_schedule_milestones

    return load_schedule_milestones(project_root)


def extract_first_date(text: str) -> str:
    match = _ISO_DATE_RE.search(text or "")
    return match.group(0) if match else ""


def _package_tokens(package_name: str) -> list[str]:
    tokens: list[str] = []
    name = (package_name or "").strip()
    if not name:
        return tokens
    tokens.append(name.lower())
    for part in re.split(r"[_\-\s/]+", name):
        part = part.strip().lower()
        if len(part) >= 2 and part not in tokens:
            tokens.append(part)
    return tokens


def _milestone_match_score(
    milestone: dict[str, str],
    keywords: list[str],
    package_tokens: list[str],
    package_name: str = "",
) -> int:
    hay = (milestone.get("name") or "").lower()
    score = 0
    for keyword in keywords:
        if keyword.lower() in hay:
            score += 10

    def _key(text: str) -> str:
        return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "", (text or "").lower())

    milestone_pkg = str(milestone.get("package") or "").strip()
    if package_name and milestone_pkg:
        pk = _key(package_name)
        mk = _key(milestone_pkg)
        if pk and mk and (pk == mk or pk in mk or mk in pk):
            score += 50
        else:
            # Different package — heavily penalize so BMS dates don't win for Glazing
            score -= 40
    for token in package_tokens:
        if token in hay:
            score += 5
    return score


def _find_best_schedule_milestone(
    milestones: list[dict[str, str]],
    keywords: list[str],
    package_name: str = "",
    *,
    min_score: int = 5,
) -> dict[str, str] | None:
    package_tokens = _package_tokens(package_name)
    best: dict[str, str] | None = None
    best_score = 0
    for milestone in milestones:
        score = _milestone_match_score(
            milestone,
            keywords,
            package_tokens,
            package_name=package_name,
        )
        if score > best_score:
            best_score = score
            best = milestone
    if best_score < min_score:
        for milestone in milestones:
            score = _milestone_match_score(milestone, keywords, [], package_name="")
            if score >= 10 and score > best_score:
                best_score = score
                best = milestone
    return best if best_score >= min_score else None


def _find_schedule_milestone(
    milestones: list[dict[str, str]],
    keywords: list[str],
) -> dict[str, str] | None:
    return _find_best_schedule_milestone(milestones, keywords, min_score=10)


def _format_schedule_window(start: str, end: str) -> str:
    start = (start or "").strip()
    end = (end or "").strip()
    if start and end and start != end:
        return f"{start} ~ {end}"
    if end:
        return f"≤ {end}"
    if start:
        return f"≤ {start}"
    return ""


def infer_schedule_window_for_payment_node(
    node: dict[str, Any],
    *,
    milestones: list[dict[str, str]],
    package_name: str = "",
    delivery_schedule: str = "",
    critical_nodes: list[dict[str, Any]] | None = None,
) -> str:
    # Prefer critical_nodes (incl. PM manual timeline edits) over stale card fields.
    from_cn = infer_window_from_critical_nodes(
        node,
        critical_nodes=critical_nodes or [],
        package_name=package_name,
    )
    if from_cn:
        return from_cn

    explicit = str(node.get("schedule_window") or "").strip()
    if explicit:
        return explicit

    haystack = " ".join(
        str(node.get(key) or "") for key in ("name", "condition", "milestone")
    )
    delivery_date = extract_first_date(delivery_schedule)
    start_milestone = _find_schedule_milestone(milestones, ["开工", "start", "commencement"])
    handover_milestone = _find_schedule_milestone(
        milestones,
        ["移交", "handover", "竣工", "completion"],
    )
    start_date = start_milestone["date"] if start_milestone else ""
    handover_date = handover_milestone["date"] if handover_milestone else ""

    for pattern, keywords, mode in _PAYMENT_SCHEDULE_RULES:
        if not re.search(pattern, haystack, re.I):
            continue

        matched = _find_best_schedule_milestone(
            milestones,
            keywords,
            package_name,
        )

        if mode == "before_start":
            if start_date:
                return f"≤ {start_date}"
            return ""

        if mode == "delivery":
            if matched:
                end = matched.get("end_date") or matched["date"]
                return _format_schedule_window(matched["date"], end)
            if delivery_date:
                return f"≤ {delivery_date}"
            return ""

        if mode == "range":
            if matched:
                end = matched.get("end_date") or matched["date"]
                if start_date and matched["date"] > start_date:
                    return _format_schedule_window(start_date, end)
                return f"≤ {end}"
            return ""

        if mode in {"before_handover", "at_or_before"}:
            if matched:
                return f"≤ {matched['date']}"
            if handover_date:
                return f"≤ {handover_date}"
            return ""

    return ""


def infer_payment_attachments(
    node: dict[str, Any],
    *,
    package_type: str = "",
) -> str:
    explicit = str(node.get("payment_attachments") or "").strip()
    if explicit:
        return explicit
    haystack = " ".join(
        str(node.get(key) or "")
        for key in ("name", "condition", "milestone")
    )
    for pattern, attachments in _PAYMENT_ATTACHMENT_RULES:
        if re.search(pattern, haystack, re.I):
            if package_type == "总包" and re.search(r"进度|progress|施工", haystack, re.I):
                return (
                    "请款单；经签认的月度对账单/进度款申请单；工程量清单；"
                    "现场进度照片；发票"
                )
            return attachments
    if package_type == "总包":
        return "请款单；经签认的支持性文件（签认单/报告/清单）；发票"
    return _DEFAULT_PAYMENT_ATTACHMENTS


def enrich_payment_nodes(
    nodes: list[dict[str, Any]],
    *,
    package_type: str = "",
    schedule_milestones: list[dict[str, str]] | None = None,
    package_name: str = "",
    delivery_schedule: str = "",
    critical_nodes: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    enriched: list[dict[str, Any]] = []
    for node in nodes:
        copy = dict(node)
        if not str(copy.get("payment_attachments") or "").strip():
            copy["payment_attachments"] = infer_payment_attachments(
                copy, package_type=package_type
            )
        if schedule_milestones is not None or critical_nodes:
            window = infer_schedule_window_for_payment_node(
                copy,
                milestones=schedule_milestones or [],
                package_name=package_name,
                delivery_schedule=delivery_schedule,
                critical_nodes=critical_nodes,
            )
            if window:
                copy["schedule_window"] = window
                if not str(copy.get("milestone") or "").strip():
                    dates = _ISO_DATE_RE.findall(window)
                    if dates:
                        copy["milestone"] = dates[-1]
        enriched.append(copy)
    return enriched


def first_pending_payment_node(payment_nodes: list[dict[str, Any]]) -> str:
    if not payment_nodes:
        return ""
    node = payment_nodes[0]
    name = node.get("name", "")
    pct = node.get("pct", "")
    return f"{name} {pct}".strip()


def registry_path() -> Path:
    import os
    if os.environ.get("PMSKILLS_REGISTRY"):
        return Path(os.environ["PMSKILLS_REGISTRY"])
    # Skill-local registry (works without writing to home); override with PMSKILLS_REGISTRY
    return Path(__file__).resolve().parent.parent / "mailbox-registry.json"


def register_project(root: Path, code: str) -> None:
    path = registry_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    data = load_json(path, {"projects": []})
    projects = data.setdefault("projects", [])
    root_s = str(root.expanduser().resolve())
    for p in projects:
        if p.get("root") == root_s or p.get("code") == code:
            p["root"] = root_s
            p["code"] = code
            save_json(path, data)
            return
    projects.append({"root": root_s, "code": code})
    save_json(path, data)


def read_mailbox_password(project_root: Path, project_code: str = "") -> str:
    """Resolve IMAP password: env → project .mailbox.secret → skill/home secret files."""
    import os

    for key in (
        f"PMSKILLS_MAIL_PASS_{project_code}" if project_code else "",
        "PMSKILLS_MAIL_PASS",
    ):
        if key and os.environ.get(key):
            return os.environ[key].strip()

    skill_secret = Path(__file__).resolve().parent.parent / ".mailbox.secret"
    for candidate in (
        project_root / ".mailbox.secret",
        skill_secret,
        Path.home() / ".pmskills" / "mailbox.secret",
    ):
        try:
            if candidate.exists():
                return candidate.read_text(encoding="utf-8").strip()
        except OSError:
            continue
    return ""


def write_mailbox_secret(project_root: Path, secret: str) -> Path:
    """Write IMAP auth code to `{project}/.mailbox.secret` (local only; never project.json).

    Also ensures `.mailbox.secret` is listed in the project root `.gitignore` when present
    or created as a one-line ignore file.
    """
    import os
    import stat

    root = project_root.expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    secret = (secret or "").strip()
    if not secret:
        raise ValueError("mailbox secret is empty")

    path = root / ".mailbox.secret"
    path.write_text(secret + "\n", encoding="utf-8")
    try:
        os.chmod(path, stat.S_IRUSR | stat.S_IWUSR)  # 0o600
    except OSError:
        pass

    gi = root / ".gitignore"
    line = ".mailbox.secret"
    if gi.is_file():
        text = gi.read_text(encoding="utf-8")
        if line not in text.splitlines():
            suffix = "" if text.endswith("\n") or not text else "\n"
            gi.write_text(text + suffix + line + "\n", encoding="utf-8")
    else:
        gi.write_text(line + "\n", encoding="utf-8")
    return path
